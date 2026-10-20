"""Shared styling constants for the petroleum engineering suite workbook."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ---------------------------------------------------------------- palette
NAVY = "1F3864"          # sheet title band
SECTION_FILL = "DDEBF7"  # section band
CALC_FILL = "F2F2F2"     # calculator title row
INPUT_FILL = "FFF2CC"    # user-editable cells (yellow)
LINK_FILL = "DAEEF3"     # cells linked to the Well Data sheet (light blue)
OUTPUT_FILL = "E2EFDA"   # computed cells (green)
SI_TEXT = "595959"       # SI mirror columns
EQ_TEXT = "2E5E8C"       # equation text color

TAB_COLORS = {
    "reference": "808080",
    "drilling": "2E75B6",
    "completions": "548235",
    "interventions": "C55A11",
}

# ---------------------------------------------------------------- fonts
F_TITLE = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
F_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="D9E2F3")
F_HEADER = Font(name="Calibri", size=10, bold=True, color="404040")
F_SECTION = Font(name="Calibri", size=12, bold=True, color=NAVY)
F_CALC = Font(name="Calibri", size=11, bold=True, color="000000")
F_EQ = Font(name="Consolas", size=10, italic=True, color=EQ_TEXT)
F_LABEL = Font(name="Calibri", size=11)
F_INPUT = Font(name="Calibri", size=11, bold=True)
F_OUTPUT = Font(name="Calibri", size=11, bold=True, color="1E4620")
F_UNIT = Font(name="Calibri", size=10, color="595959")
F_SI = Font(name="Calibri", size=10, color=SI_TEXT)
F_NOTE = Font(name="Calibri", size=9, italic=True, color="7F7F7F")

FILL_TITLE = PatternFill("solid", fgColor=NAVY)
FILL_SECTION = PatternFill("solid", fgColor=SECTION_FILL)
FILL_CALC = PatternFill("solid", fgColor=CALC_FILL)
FILL_INPUT = PatternFill("solid", fgColor=INPUT_FILL)
FILL_LINK = PatternFill("solid", fgColor=LINK_FILL)
F_LINK = Font(name="Calibri", size=11, bold=True, color="1F4E79")
FILL_OUTPUT = PatternFill("solid", fgColor=OUTPUT_FILL)

_thin = Side(style="thin", color="BFBFBF")
BORDER_BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

AL_LEFT = Alignment(horizontal="left", vertical="center")
AL_RIGHT = Alignment(horizontal="right", vertical="center")
AL_CENTER = Alignment(horizontal="center", vertical="center")
AL_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)

COL_WIDTHS = {"A": 42, "B": 14, "C": 11, "D": 14, "E": 11, "F": 78}
