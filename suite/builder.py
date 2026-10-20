"""SheetBuilder: lays out calculator blocks on a worksheet.

Layout per sheet:
  row 1  title band (merged A:F)
  row 2  subtitle band (merged A:F)
  row 4  column headers: Parameter | Value | Unit | SI Value | SI Unit | Notes / Equation
  row 5+ sections and calculator blocks (freeze panes at A5)

Each calculator block:
  - bold title row on a gray band, equation text (Consolas, blue) in the Notes column
  - input rows   (yellow value cell)  -> c.inp() returns "$B$n" for use in formulas
  - output rows  (green value cell, live Excel formula) -> c.out() also returns "$B$n"
  - SI mirror in columns D:E, converted live from column B
"""
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

UNLOCKED = Protection(locked=False)

from . import style as S
from .convert import CONV

DEFAULT_FMT = "#,##0.00"


class Calc:
    def __init__(self, sb: "SheetBuilder"):
        self.sb = sb

    def _si_cells(self, row, unit, si, fmt):
        ws = self.sb.ws
        if si is None and unit in CONV:
            si = CONV[unit]
        if si is None:
            return
        # IFERROR->NA keeps SI mirrors quiet when the source cell is an
        # intentional NA() (e.g. no-solution outputs)
        if isinstance(si, tuple) and callable(si[1]):
            label, fn = si
            inner = fn(f"$B${row}").lstrip("=")
            ws.cell(row=row, column=4, value=f"=IFERROR({inner},NA())")
        else:
            label, factor = si
            ws.cell(row=row, column=4, value=f"=IFERROR($B${row}*{factor},NA())")
        d = ws.cell(row=row, column=4)
        d.font = S.F_SI
        d.number_format = fmt
        d.alignment = S.AL_RIGHT
        e = ws.cell(row=row, column=5, value=label)
        e.font = S.F_SI
        e.alignment = S.AL_LEFT

    def _row(self, label, value, unit, si, note, fmt, is_input):
        sb = self.sb
        ws = sb.ws
        r = sb.row
        a = ws.cell(row=r, column=1, value=label)
        a.font = S.F_LABEL
        a.alignment = S.AL_LEFT
        b = ws.cell(row=r, column=2, value=value)
        b.number_format = fmt
        b.alignment = S.AL_RIGHT
        b.border = S.BORDER_BOX
        if is_input == "link":
            b.fill = S.FILL_LINK
            b.font = S.F_LINK
            b.protection = UNLOCKED
        elif is_input:
            b.fill = S.FILL_INPUT
            b.font = S.F_INPUT
            b.protection = UNLOCKED
        else:
            b.fill = S.FILL_OUTPUT
            b.font = S.F_OUTPUT
        c = ws.cell(row=r, column=3, value=unit or "")
        c.font = S.F_UNIT
        c.alignment = S.AL_LEFT
        self._si_cells(r, unit, si, fmt)
        if note:
            f = ws.cell(row=r, column=6, value=note)
            f.font = S.F_NOTE
            f.alignment = S.AL_WRAP
        sb.row += 1
        return f"$B${r}"

    def inp(self, label, value, unit=None, si=None, note="", fmt=DEFAULT_FMT):
        """Yellow user-input cell. Returns its absolute reference."""
        return self._row(label, value, unit, si, note, fmt, True)

    def link(self, label, defined_name, unit=None, si=None, note="", fmt=DEFAULT_FMT):
        """Light-blue cell linked to a Well Data defined name (e.g. "WD_TVD").
        Unlocked so the user can overwrite it locally. Returns its reference."""
        note = note or "Linked to Well Data sheet - overwrite locally if this well differs."
        return self._row(label, f"={defined_name}", unit, si, note, fmt, "link")

    def out(self, label, formula, unit=None, si=None, note="", fmt=DEFAULT_FMT):
        """Green computed cell. `formula` is the Excel formula without '='.
        Returns its absolute reference."""
        return self._row(label, "=" + formula, unit, si, note, fmt, False)

    def txt(self, label, text=""):
        """Informational row (no value cell)."""
        sb = self.sb
        a = sb.ws.cell(row=sb.row, column=1, value=label)
        a.font = S.F_NOTE
        a.alignment = S.AL_WRAP
        if text:
            f = sb.ws.cell(row=sb.row, column=6, value=text)
            f.font = S.F_NOTE
            f.alignment = S.AL_WRAP
        sb.row += 1


class SheetBuilder:
    registry = []  # (sheet_name, group, description) in build order, for the cover TOC

    def __init__(self, wb, name, group, title, subtitle="", description=""):
        assert len(name) <= 31, name
        self.ws: Worksheet = wb.create_sheet(name)
        self.ws.sheet_properties.tabColor = S.TAB_COLORS[group]
        SheetBuilder.registry.append((name, group, description or subtitle))
        for col, w in S.COL_WIDTHS.items():
            self.ws.column_dimensions[col].width = w

        self.ws.merge_cells("A1:F1")
        t = self.ws.cell(row=1, column=1, value=title)
        t.font = S.F_TITLE
        t.fill = S.FILL_TITLE
        t.alignment = S.AL_LEFT
        self.ws.row_dimensions[1].height = 26
        self.ws.merge_cells("A2:F2")
        st = self.ws.cell(row=2, column=1, value=subtitle)
        st.font = S.F_SUBTITLE
        st.fill = S.FILL_TITLE
        st.alignment = S.AL_LEFT
        for col in range(1, 7):
            self.ws.cell(row=1, column=col).fill = S.FILL_TITLE
            self.ws.cell(row=2, column=col).fill = S.FILL_TITLE

        headers = ["Parameter", "Value", "Unit", "SI Value", "SI Unit", "Notes / Equation"]
        for i, h in enumerate(headers, start=1):
            c = self.ws.cell(row=4, column=i, value=h)
            c.font = S.F_HEADER
            c.alignment = S.AL_LEFT if i in (1, 6) else S.AL_RIGHT if i in (2, 4) else S.AL_LEFT
        self.ws.freeze_panes = "A5"
        self.row = 5

    def section(self, text):
        self.row += 1
        r = self.row
        self.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c = self.ws.cell(row=r, column=1, value=text)
        c.font = S.F_SECTION
        c.fill = S.FILL_SECTION
        c.alignment = S.AL_LEFT
        for col in range(1, 7):
            self.ws.cell(row=r, column=col).fill = S.FILL_SECTION
        self.ws.row_dimensions[r].height = 20
        self.row += 1
        return self

    def calc(self, name, eq=None, note=None):
        """Start a calculator block; equation text goes in the Notes column."""
        self.row += 1
        r = self.row
        c = self.ws.cell(row=r, column=1, value=name)
        c.font = S.F_CALC
        c.fill = S.FILL_CALC
        c.alignment = S.AL_LEFT
        for col in range(1, 7):
            self.ws.cell(row=r, column=col).fill = S.FILL_CALC
        if eq:
            f = self.ws.cell(row=r, column=6, value=eq)
            f.font = S.F_EQ
            f.alignment = S.AL_WRAP
        self.row += 1
        if note:
            n = self.ws.cell(row=self.row, column=6, value=note)
            n.font = S.F_NOTE
            n.alignment = S.AL_WRAP
            a = self.ws.cell(row=self.row, column=1, value="")
            self.row += 1
        return Calc(self)
