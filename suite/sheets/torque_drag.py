"""Drilling: torque & drag - Johancsik soft-string model over the survey.

Per survey segment (station r to r+1, properties stored on the upper node r):

    N  = sqrt( (T x dAzi x sin(th))^2 + (T x dInc + w_b x dMD x sin(th))^2 )
    dT = w_b x dMD x cos(th)  +/- mu x N      (+ pick-up, - slack-off)
    dM = mu x N x r_tj                        (rotating torque)

with T taken at the lower node (marching bottom-up), th the segment average
inclination, w_b the buoyed weight per foot. Buckling screened per segment
with Dawson-Paslay sinusoidal limit  F_sin = 2 x sqrt(E I w_b sin(th) / r).

Soft-string: pipe stiffness and stabilizer standoff ignored - adequate for
DP-dominated strings; calibrate friction factors to measured hookloads.
"""
from openpyxl.chart import Reference, ScatterChart, Series

from .. import style as S
from ..builder import UNLOCKED, SheetBuilder
from .directional_survey import DATA_FIRST as SV_FIRST
from .directional_survey import DATA_LAST as SV_LAST

SV = "'Directional Survey'"
N_ROWS = SV_LAST - SV_FIRST + 1


def build(wb):
    sb = SheetBuilder(
        wb, "Torque & Drag", "drilling",
        "Torque & Drag - Soft-String (Johancsik)",
        "Pick-up / slack-off / rotating tension and torque along the survey, hookloads, "
        "buckling screen, tension & torque plots",
        description="Johancsik soft-string T&D over the Directional Survey: hookloads, "
                    "surface torque, Dawson-Paslay buckling flags, profile charts",
    )
    ws = sb.ws

    sb.section("Inputs (light blue = linked to Well Data)")
    c = sb.calc("String & Friction",
                note="Calibrate friction factors to measured hookloads when available. "
                     "Soft-string model: DP-dominated strings; stiff BHAs read high.")
    mw = c.link("Mud weight", "WD_OMW", "ppg")
    mdbit = c.link("Bit / TD measured depth", "WD_MD", "ft", fmt="#,##0")
    shoe = c.link("Casing shoe MD", "WD_ShoeMD", "ft", fmt="#,##0")
    hole = c.link("Open hole diameter", "WD_HoleDiam", "in", fmt="0.000")
    csgid = c.link("Casing ID", "WD_CsgID", "in", fmt="0.000")
    dpwt = c.link("Drill pipe adjusted weight", "WD_DPWt", "lb/ft", fmt="0.00")
    bhawt = c.link("BHA weight per foot", "WD_BHAWt", "lb/ft", fmt="0.00")
    bhalen = c.link("BHA length", "WD_BHALen", "ft", fmt="#,##0")
    mucsd = c.link("Friction factor - cased", "WD_MuCased", None, fmt="0.00")
    muoh = c.link("Friction factor - open hole", "WD_MuOpen", None, fmt="0.00")
    wob = c.link("Weight on bit (slack-off / rotating case)", "WD_WOB", "klbf", fmt="0.0")
    bittq = c.link("Bit torque (rotating case)", "WD_BitTq", "ft-lbf", fmt="#,##0")
    tjod = c.link("Tool joint OD (torque radius)", "WD_TJOD", "in", fmt="0.000")
    block = c.link("Block + top drive weight", "WD_BlockWt", "klbf", fmt="#,##0")
    dpod = c.inp("Drill pipe body OD (buckling)", 5.0, "in", fmt="0.000")
    dpid = c.inp("Drill pipe body ID (buckling)", 4.276, "in", fmt="0.000")
    emod = c.inp("Steel modulus (E)", 30000000, "psi", fmt="#,##0",
                 note="Api RP 7G commonly 29.6-30 x 10^6 psi.")
    bf = c.out("Buoyancy factor", f"(65.5-{mw})/65.5", None, fmt="0.000")
    iner = c.out("DP moment of inertia (I)", f"PI()/64*({dpod}^4-{dpid}^4)", None,
                 fmt="0.00", note="in^4")

    # ------------------------------------------------------------ station table
    sb.section("Station Table (marches bottom-up from the bit; rows mirror the survey)")
    ws.cell(row=sb.row, column=1,
            value="Tension in lbf at each survey node. Pick-up boundary: 0 at bit "
                  "(off-bottom); slack-off & rotating boundary: -WOB. Torque boundary: bit "
                  "torque. Buckling screen on the slack-off case.").font = S.F_NOTE
    ws.merge_cells(start_row=sb.row, start_column=1, end_row=sb.row, end_column=6)
    ws.cell(row=sb.row, column=1).alignment = S.AL_WRAP
    ws.row_dimensions[sb.row].height = 30
    sb.row += 1
    hdr = sb.row
    heads = ["MD (ft)", "Inc (deg)", "dMD (ft)", "th avg (deg)", "dInc (rad)", "dAzi (rad)",
             "w_b (lb/ft)", "T pick-up (lbf)", "T slack-off (lbf)", "T rotating (lbf)",
             "Torque (ft-lbf)", "F_sin limit (lbf)", "Buckling flag",
             "MD_c*", "PU_c*", "SO_c*", "ROT_c*", "TQ_c*"]
    for i, h in enumerate(heads, start=1):
        cell = ws.cell(row=hdr, column=i, value=h)
        cell.font = S.F_HEADER
        cell.border = S.BORDER_BOX
        cell.alignment = S.AL_RIGHT
    for col, w in zip("ABCDEFGHIJKLMNOPQR", (10, 8, 8, 9, 9, 9, 9, 13, 13, 13, 12, 12, 22,
                                             9, 9, 9, 9, 9)):
        ws.column_dimensions[col].width = w

    tf = hdr + 1                      # first table row
    tl = tf + N_ROWS - 1              # last table row
    for i in range(N_ROWS):
        r = tf + i
        n = r + 1                     # next (deeper) table row
        sv = SV_FIRST + i             # matching survey row
        svn = sv + 1
        A, B, C, D, E, F, G = (f"$A{r}", f"$B{r}", f"$C{r}", f"$D{r}", f"$E{r}",
                               f"$F{r}", f"$G{r}")
        ws.cell(row=r, column=1, value=f'=IF({SV}!A{sv}="","",{SV}!A{sv})')
        ws.cell(row=r, column=2, value=f'=IF({A}="","",{SV}!B{sv})')
        if i < N_ROWS - 1:
            nxt_md = f"{SV}!A{svn}"
            ws.cell(row=r, column=3, value=f'=IF(OR({A}="",{nxt_md}=""),"",{nxt_md}-{A})')
            ws.cell(row=r, column=4, value=f'=IF({C}="","",({B}+{SV}!B{svn})/2)')
            ws.cell(row=r, column=5, value=f'=IF({C}="","",RADIANS({SV}!B{svn}-{B}))')
            ws.cell(row=r, column=6,
                    value=f'=IF({C}="","",RADIANS({SV}!C{svn}-{SV}!C{sv}))')
        else:
            ws.cell(row=r, column=3, value='=""')
        ws.cell(row=r, column=7,
                value=f'=IF({C}="","",{bf}*IF(({A}+{A}+{C})/2>{mdbit}-{bhalen},{bhawt},{dpwt}))')
        mu = f"IF(({A}+{C}/2)<={shoe},{mucsd},{muoh})"
        sth = f"SIN(RADIANS({D}))"
        cth = f"COS(RADIANS({D}))"
        for col, sign, letter, bound in ((8, "+", "H", "0"),
                                         (9, "-", "I", f"-{wob}*1000"),
                                         (10, None, "J", f"-{wob}*1000")):
            tnext = f"${letter}{n}"
            norm = (f"SQRT(({tnext}*{F}*{sth})^2+({tnext}*{E}+{G}*{C}*{sth})^2)")
            if sign is None:  # rotating: no axial friction
                frm = (f'=IF({A}="","",IF({C}="",{bound},'
                       f"{tnext}+{G}*{C}*{cth}))")
            else:
                frm = (f'=IF({A}="","",IF({C}="",{bound},'
                       f"{tnext}+{G}*{C}*{cth}{sign}{mu}*{norm}))")
            ws.cell(row=r, column=col, value=frm)
        # torque: friction normal force from the rotating tension
        tnext = f"$J{n}"
        norm = f"SQRT(({tnext}*{F}*{sth})^2+({tnext}*{E}+{G}*{C}*{sth})^2)"
        ws.cell(row=r, column=11,
                value=f'=IF({A}="","",IF({C}="",{bittq},$K{n}+{mu}*{norm}*{tjod}/24))')
        # Dawson-Paslay sinusoidal limit; 3-deg floor avoids the vertical singularity
        rcl = f"(IF({A}<={shoe},{csgid},{hole})-{dpod})/2"
        ws.cell(row=r, column=12,
                value=f'=IF({C}="","",2*SQRT({emod}*{iner}*({G}/12)*'
                      f"SIN(RADIANS(MAX({D},3)))/{rcl}))")
        ws.cell(row=r, column=13,
                value=f'=IF({C}="","",IF($I{r}<-1.41*$L{r},"HELICAL BUCKLING - CHECK",'
                      f'IF($I{r}<-$L{r},"SINUSOIDAL - MARGINAL","")))')
        # numeric chart helpers (klbf / kft-lbf)
        ws.cell(row=r, column=14, value=f'=IF({A}="",NA(),{A})')
        for col, src in ((15, f"$H{r}"), (16, f"$I{r}"), (17, f"$J{r}")):
            ws.cell(row=r, column=col, value=f'=IF({A}="",NA(),{src}/1000)')
        ws.cell(row=r, column=18, value=f'=IF({A}="",NA(),$K{r}/1000)')
        for col in range(1, 13):
            cell = ws.cell(row=r, column=col)
            cell.number_format = "#,##0" if col in (1, 3, 8, 9, 10, 11, 12) else "0.000"
            cell.alignment = S.AL_RIGHT
        ws.cell(row=r, column=13).font = S.F_NOTE
        for col in range(14, 19):
            cc = ws.cell(row=r, column=col)
            cc.font = S.F_NOTE
            cc.number_format = "0.0"
    sb.row = tl + 1

    # ------------------------------------------------------------ surface results
    sb.section("Surface Results")
    c = sb.calc("Hookloads & Torque",
                eq="Hookload = string tension at surface + block weight",
                note="Compare against measured pick-up/slack-off to calibrate friction "
                     "factors. Overpull margin vs pipe yield is on the Drill String sheet.")
    c.out("Pick-up hookload", f"$H${tf}/1000+{block}", "klbf", fmt="0.0")
    c.out("Slack-off hookload", f"$I${tf}/1000+{block}", "klbf", fmt="0.0")
    c.out("Rotating on-bottom hookload", f"$J${tf}/1000+{block}", "klbf", fmt="0.0",
          note="Rotating case carries WOB at the bit; set WOB to 0 on Well Data for "
               "rotating off-bottom.")
    c.out("Surface torque (rotating)", f"$K${tf}", "ft-lbf", fmt="#,##0")
    c.out("Buckling anywhere in string",
          f'IF(COUNTIF($M${tf}:$M${tl},"*BUCKLING*")>0,"CHECK: helical buckling flagged - '
          f'reduce WOB or add HWDP",IF(COUNTIF($M${tf}:$M${tl},"*SINUSOIDAL*")>0,'
          f'"MARGINAL - sinusoidal buckling flagged","OK - no buckling flags"))',
          None, fmt="@")

    # ------------------------------------------------------------ charts
    tens = ScatterChart()
    tens.varyColors = False
    tens.title = "Tension vs MD"
    tens.x_axis.title = "Tension (klbf)"
    tens.y_axis.title = "MD (ft)"
    tens.y_axis.scaling.orientation = "maxMin"
    for ax in (tens.x_axis, tens.y_axis):
        ax.delete = False
        ax.tickLblPos = "nextTo"
    tens.height = 12
    tens.width = 16
    yref = Reference(ws, min_col=14, min_row=tf, max_row=tl)
    for col, name in ((15, "Pick-up"), (16, "Slack-off"), (17, "Rotating")):
        xref = Reference(ws, min_col=col, min_row=tf, max_row=tl)
        s = Series(yref, xref, title=name)
        tens.series.append(s)
    ws.add_chart(tens, f"T{hdr}")

    tq = ScatterChart()
    tq.varyColors = False
    tq.title = "Torque vs MD"
    tq.x_axis.title = "Torque (kft-lbf)"
    tq.y_axis.title = "MD (ft)"
    tq.y_axis.scaling.orientation = "maxMin"
    for ax in (tq.x_axis, tq.y_axis):
        ax.delete = False
        ax.tickLblPos = "nextTo"
    tq.height = 12
    tq.width = 16
    xref = Reference(ws, min_col=18, min_row=tf, max_row=tl)
    s = Series(yref, xref, title="Rotating torque")
    tq.series.append(s)
    ws.add_chart(tq, f"T{hdr + 25}")
