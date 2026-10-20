"""Drilling: maximum trip speed vs depth from surge/swab limits.

Burkhardt closed-ended clinging model with laminar Bingham-plastic friction
(same formulation as the Swab & Surge calcs on the Hydraulics and
Tubing & Packers sheets), inverted to solve for pipe speed:

    Ve = Vp x (0.45 + Dp^2/(Dc^2 - Dp^2)) = Vp x K
    dP = L x PV x Ve / (60000 x (Dc-Dp)^2) + L x YP / (200 x (Dc-Dp))
    Vp_max = (dP_allow - L x YP/(200 x (Dc-Dp))) x 60000 x (Dc-Dp)^2 / (L x PV x K)

Schedule table geometry:
  A=bit depth (yellow, MD=TVD assumed)  B=allow surge psi  C=allow swab psi
  D=max run speed  E=max pull speed  F=s/stand run  G=s/stand pull
  H..J = NA()-guarded numeric helpers for the speed chart
"""
from openpyxl.chart import Reference, ScatterChart, Series

from .. import style as S
from ..builder import UNLOCKED, SheetBuilder

N_ROWS = 12
STAND_FT = 93


def _cell(ws, row, col, value, fmt="#,##0.00", is_input=False, is_output=False):
    c = ws.cell(row=row, column=col, value=value)
    c.number_format = fmt
    c.border = S.BORDER_BOX
    c.alignment = S.AL_RIGHT
    if is_input:
        c.fill = S.FILL_INPUT
        c.font = S.F_INPUT
        c.protection = UNLOCKED
    elif is_output:
        c.fill = S.FILL_OUTPUT
        c.font = S.F_OUTPUT
    else:
        c.font = S.F_LABEL
    return c


def build(wb):
    sb = SheetBuilder(
        wb, "Trip Speed Schedule", "drilling",
        "Maximum Trip Speed vs Depth - Surge & Swab Limits",
        "Burkhardt closed-ended clinging + laminar Bingham friction, inverted for pipe "
        "speed; running and pulling speed schedule with seconds per stand",
        description="Max running/pulling speed vs bit depth from surge (frac/loss) and "
                    "swab (pore/kick) EMW limits, seconds per 93-ft stand, speed chart",
    )
    ws = sb.ws
    for col, w in {"G": 12, "H": 10, "I": 10, "J": 10}.items():
        ws.column_dimensions[col].width = w

    sb.section("Geometry & Fluid")

    c = sb.calc("Annulus Geometry",
                note="Closed-ended (float / bullnose / plugged bit) is the worst case and "
                     "is what the schedule below assumes. Open-ended pipe surges less.")
    dc = c.link("Casing / hole ID (Dc)", "WD_CsgID", "in", fmt="0.000",
                note="Linked to Well Data. Use the smallest restriction the pipe runs "
                     "through (casing ID here; overwrite with open-hole size if smaller).")
    dp = c.inp("Pipe OD (Dp)", 5.0, "in", fmt="0.000")
    c.inp("Pipe ID", 4.276, "in", fmt="0.000",
          note="Reference only - closed-ended assumption ignores the pipe bore.")
    kf = c.out("Clinging factor (K)", f"0.45+{dp}^2/({dc}^2-{dp}^2)", None, fmt="0.000",
               note="Burkhardt closed-ended: Ve = Vp x K, K = 0.45 + Dp^2/(Dc^2-Dp^2).")

    c = sb.calc("Mud Properties")
    mw = c.link("Mud weight (MW)", "WD_OMW", "ppg")
    pv = c.link("Plastic viscosity (PV)", "WD_PV", None, fmt="0.0", note="cP.")
    yp = c.link("Yield point (YP)", "WD_YP", None, fmt="0.0", note="lb/100 ft2.")

    sb.section("Pressure Limits")

    c = sb.calc("Allowable Surge & Swab",
                eq="dP_surge_allow = (EMW_frac - MW) x 0.052 x TVD;   "
                   "dP_swab_allow = (MW - EMW_pore) x 0.052 x TVD",
                note="Surge side is limited by the frac/loss EMW (running in raises BHP); "
                     "swab side by the pore + kick-margin EMW (pulling out lowers BHP). "
                     "The table below evaluates both at each bit depth with L = depth.")
    fracemw = c.link("Frac / loss limit EMW", "WD_LOT", "ppg",
                     note="Linked to the LOT/FIT EMW - overwrite with a loss-zone EMW if lower.")
    poreemw = c.inp("Pore / swab limit EMW", 12.0, "ppg", fmt="0.00",
                    note="Pore pressure EMW plus kick margin on the swab side. If this is "
                         "at/above MW, tripping out swabs the well in - raise MW first.")
    tvd = c.link("TVD at TD (reference)", "WD_TVD", "ft", fmt="#,##0")
    vcap = c.inp("Rig handling speed cap", 400, "ft/min", fmt="#,##0",
                 note="Practical block-speed/handling ceiling; the schedule shows "
                      "MIN(cap, hydraulic limit). Laminar surge model loses validity at "
                      "very high speeds anyway.")
    smarg = c.out("Surge margin (frac EMW - MW)", f"{fracemw}-{mw}", "ppg", fmt="0.00")
    wmarg = c.out("Swab margin (MW - swab limit)", f"{mw}-{poreemw}", "ppg", fmt="0.00")
    c.out("Surge status",
          f'IF({smarg}>=0.5,"OK - surge margin adequate",'
          f'IF({smarg}>0,"MARGINAL - thin surge margin - run slowly",'
          f'"DO NOT run in - MW at/above frac limit"))',
          None, fmt="@")
    c.out("Swab status",
          f'IF({wmarg}>=0.3,"OK - swab margin adequate",'
          f'IF({wmarg}>0,"MARGINAL - thin swab margin - pull slowly",'
          f'"DO NOT trip out on this mud - MW at/below swab limit; raise MW first"))',
          None, fmt="@")

    # ------------------------------------------------------------ schedule table
    sb.row += 1
    band_row = sb.row
    ws.merge_cells(start_row=band_row, start_column=1, end_row=band_row, end_column=6)
    band = ws.cell(row=band_row, column=1,
                   value="Trip Speed Schedule - enter bit depths (yellow, MD = TVD assumed; "
                         "use survey TVDs on deviated wells)")
    band.font = S.F_SECTION
    band.fill = S.FILL_SECTION
    for col in range(1, 7):
        ws.cell(row=band_row, column=col).fill = S.FILL_SECTION
    sb.row += 1

    hdr_row = sb.row
    headers = ["Bit depth (ft)", "Allow surge (psi)", "Allow swab (psi)",
               "Max run (ft/min)", "Max pull (ft/min)", "s/stand run", "s/stand pull",
               "RUNh*", "PULLh*", "DEPTHh*"]
    for i, h in enumerate(headers, start=1):
        hc = ws.cell(row=hdr_row, column=i, value=h)
        hc.font = S.F_HEADER
        hc.border = S.BORDER_BOX
        hc.alignment = S.AL_RIGHT if i > 1 else S.AL_LEFT
    sb.row += 1

    first = sb.row
    last = first + N_ROWS - 1
    yp_term = f"$A{{r}}*{yp}/(200*({dc}-{dp}))"
    for i in range(N_ROWS):
        r = first + i
        ypt = yp_term.format(r=r)
        _cell(ws, r, 1, 1000 * (i + 1), fmt="#,##0", is_input=True)
        _cell(ws, r, 2, f'=IF($A{r}="","",MAX(0,({fracemw}-{mw})*0.052*$A{r}))', fmt="#,##0")
        _cell(ws, r, 3, f'=IF($A{r}="","",MAX(0,({mw}-{poreemw})*0.052*$A{r}))', fmt="#,##0")
        _cell(ws, r, 4,
              f'=IF($A{r}="","",MIN({vcap},MAX(0,(B{r}-{ypt})*60000*({dc}-{dp})^2'
              f"/($A{r}*{pv}*{kf}))))",
              fmt="#,##0")
        _cell(ws, r, 5,
              f'=IF($A{r}="","",MIN({vcap},MAX(0,(C{r}-{ypt})*60000*({dc}-{dp})^2'
              f"/($A{r}*{pv}*{kf}))))",
              fmt="#,##0")
        _cell(ws, r, 6, f'=IF($A{r}="","",IF(D{r}<=0,"-",{STAND_FT}/D{r}*60))', fmt="#,##0")
        _cell(ws, r, 7, f'=IF($A{r}="","",IF(E{r}<=0,"-",{STAND_FT}/E{r}*60))', fmt="#,##0")
        # NA()-guarded numeric helpers for the chart
        helpers = (
            (8, f'=IF($A{r}="",NA(),D{r})'),
            (9, f'=IF($A{r}="",NA(),E{r})'),
            (10, f'=IF($A{r}="",NA(),$A{r})'),
        )
        for col, formula in helpers:
            hcell = ws.cell(row=r, column=col, value=formula)
            hcell.font = S.F_NOTE
            hcell.number_format = "#,##0"
            hcell.alignment = S.AL_RIGHT
    sb.row = last + 1
    for text in (
        "* helper columns H-J feed the speed chart - do not edit. #N/A on unused rows "
        "by design (chart gaps).",
        "Crews should round speeds DOWN and time the fastest stand, not the average - "
        "peak speed within a stand can be 1.5-2x the average.",
        "Closed-ended (float/bullnose) is shown - worst case. Open-ended pipe displaces "
        "less fluid and allows faster tripping; a speed of 0 or '-' means the margin "
        "is gone at that depth on this mud.",
    ):
        note = ws.cell(row=sb.row, column=1, value=text)
        note.font = S.F_NOTE
        sb.row += 1

    # ------------------------------------------------------------ chart
    chart = ScatterChart()
    chart.varyColors = False
    chart.title = "Maximum Trip Speed vs Depth"
    chart.x_axis.title = "Pipe speed (ft/min)"
    chart.y_axis.title = "Bit depth (ft)"
    chart.y_axis.scaling.orientation = "maxMin"
    for ax in (chart.x_axis, chart.y_axis):
        ax.delete = False
        ax.tickLblPos = "nextTo"
    chart.height = 12
    chart.width = 16
    yref = Reference(ws, min_col=10, min_row=first, max_row=last)
    for col, name in ((8, "Run in (surge limit)"), (9, "Pull out (swab limit)")):
        xref = Reference(ws, min_col=col, min_row=first, max_row=last)
        s = Series(yref, xref, title=name)
        s.marker.symbol = "circle"
        s.marker.size = 4
        chart.series.append(s)
    ws.add_chart(chart, "L6")
