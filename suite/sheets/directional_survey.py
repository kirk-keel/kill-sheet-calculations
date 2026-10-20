"""Drilling: directional survey table (minimum curvature) with trajectory charts.

Paste MD / Inc / Azi stations into the yellow columns (row 1 = tie-in).
Every row computes TVD, N/S, E/W, vertical section and DLS by minimum
curvature. Helper columns I/J (MDh/TVDh) give other sheets (Kill Sheet)
a clean numeric range for MD -> TVD interpolation.

Table geometry (used by kill_sheet.py):
  headers row 6, data rows 7..106
  A=MD  B=Inc  C=Azi  D=TVD  E=VS  F=N/S  G=E/W  H=DLS  I=MDh  J=TVDh  K=DL  L=RF
"""
from openpyxl.chart import Reference, ScatterChart, Series

from .. import style as S
from ..builder import UNLOCKED, SheetBuilder

SHEET = "Directional Survey"
DATA_FIRST = 7
DATA_LAST = 106
MD_HELPER = f"'{SHEET}'!$I${DATA_FIRST}:$I${DATA_LAST}"
TVD_HELPER = f"'{SHEET}'!$J${DATA_FIRST}:$J${DATA_LAST}"
BIG = "9.9E+99"

# default example: horizontal well, KOP 5000 ft, 10 deg/100ft build to 90 deg
# at 5900 ft MD (TVD ~5573), 4600 ft lateral to TD 10500 ft MD
DEFAULT_SURVEY = (
    [(0, 0, 0), (1000, 0, 0), (2000, 0, 0), (3000, 0, 0), (4000, 0, 0),
     (4500, 0, 0), (5000, 0, 0)]
    + [(5000 + 100 * i, 10 * i, 90) for i in range(1, 10)]
    + [(6000, 90, 90), (6500, 90, 90), (7000, 90, 90), (7500, 90, 90),
       (8000, 90, 90), (8500, 90, 90), (9000, 90, 90), (9500, 90, 90),
       (10000, 90, 90), (10500, 90, 90)]
)


def interp_tvd(md_ref):
    """Excel formula (no '='): TVD at md_ref, linear interp on helper columns."""
    m = f"MATCH({md_ref},{MD_HELPER},1)"
    i0 = f"INDEX({MD_HELPER},{m})"
    t0 = f"INDEX({TVD_HELPER},{m})"
    i1 = f"IFERROR(INDEX({MD_HELPER},{m}+1),{BIG})"
    t1 = f"IFERROR(INDEX({TVD_HELPER},{m}+1),{t0})"
    return f"{t0}+({md_ref}-{i0})*({t1}-{t0})/({i1}-{i0})"


def build(wb):
    sb = SheetBuilder(
        wb, SHEET, "drilling",
        "Directional Survey - Minimum Curvature",
        "Paste MD / Inc / Azi stations; TVD, N/E, VS and DLS compute per row. "
        "Charts and the Kill Sheet schedule update automatically.",
        description="Survey upload table (100 stations, minimum curvature), TVD/NS/EW/VS/DLS, "
                    "profile & plan charts, MD->TVD interpolator",
    )
    ws = sb.ws
    ws.freeze_panes = "A7"
    for col in range(1, 7):  # clear the standard 6-column headers; this sheet is a table
        ws.cell(row=4, column=col, value=None)
    for col, w in {"A": 12, "B": 9, "C": 9, "D": 12, "E": 12, "F": 11, "G": 11,
                   "H": 13, "I": 11, "J": 11, "K": 9, "L": 9,
                   "M": 10, "N": 10, "O": 10, "P": 10}.items():
        ws.column_dimensions[col].width = w

    # row 3: VS direction input + summary outputs
    ws.cell(row=3, column=1, value="Vertical section direction").font = S.F_LABEL
    b3 = ws.cell(row=3, column=2, value=90)
    b3.fill = S.FILL_INPUT
    b3.border = S.BORDER_BOX
    b3.font = S.F_INPUT
    b3.protection = UNLOCKED
    b3.number_format = "0.0"
    ws.cell(row=3, column=3, value="deg").font = S.F_UNIT
    ws.cell(row=3, column=4, value="TVD at TD").font = S.F_LABEL
    e3 = ws.cell(row=3, column=5, value=f"=LOOKUP(9.99E+307,D{DATA_FIRST}:D{DATA_LAST})")
    e3.fill = S.FILL_OUTPUT
    e3.font = S.F_OUTPUT
    e3.border = S.BORDER_BOX
    e3.number_format = "#,##0.0"
    ws.cell(row=3, column=6, value="ft").font = S.F_UNIT
    ws.cell(row=3, column=7, value="Max DLS").font = S.F_LABEL
    h3 = ws.cell(row=3, column=8, value=f"=MAX(H{DATA_FIRST + 1}:H{DATA_LAST})")
    h3.fill = S.FILL_OUTPUT
    h3.font = S.F_OUTPUT
    h3.border = S.BORDER_BOX
    h3.number_format = "0.00"
    ws.cell(row=3, column=9, value="deg/100ft").font = S.F_UNIT

    # row 5: band
    ws.merge_cells(start_row=5, start_column=1, end_row=5, end_column=16)
    band = ws.cell(row=5, column=1,
                   value="Survey Stations - enter MD / Inc / Azi (yellow). Row 1 is the tie-in "
                         "(TVD/NS/EW also inputs there). MD must strictly increase; no blank rows "
                         "between stations; last station must reach bit/TD measured depth.")
    band.font = S.F_SECTION
    band.fill = S.FILL_SECTION
    for col in range(1, 17):
        ws.cell(row=5, column=col).fill = S.FILL_SECTION

    headers = ["MD (ft)", "Inc (deg)", "Azi (deg)", "TVD (ft)", "VS (ft)", "N/S (ft)",
               "E/W (ft)", "DLS (deg/100ft)", "MDh*", "TVDh*", "DL*", "RF*",
               "VS_c*", "TVD_c*", "EW_c*", "NS_c*"]
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=6, column=i, value=h)
        c.font = S.F_HEADER
        c.border = S.BORDER_BOX
        c.alignment = S.AL_RIGHT

    for idx in range(DATA_FIRST, DATA_LAST + 1):
        r, p = idx, idx - 1
        di = idx - DATA_FIRST  # data index
        md, inc, azi = DEFAULT_SURVEY[di] if di < len(DEFAULT_SURVEY) else (None, None, None)
        for col, val in ((1, md), (2, inc), (3, azi)):
            cell = ws.cell(row=r, column=col, value=val)
            cell.fill = S.FILL_INPUT
            cell.border = S.BORDER_BOX
            cell.font = S.F_INPUT
            cell.protection = UNLOCKED
            cell.number_format = "#,##0" if col == 1 else "0.0"
            cell.alignment = S.AL_RIGHT
        if idx == DATA_FIRST:
            # tie-in row: TVD, N/S, E/W are inputs
            for col in (4, 6, 7):
                cell = ws.cell(row=r, column=col, value=0)
                cell.fill = S.FILL_INPUT
                cell.border = S.BORDER_BOX
                cell.font = S.F_INPUT
                cell.protection = UNLOCKED
                cell.number_format = "#,##0.0"
                cell.alignment = S.AL_RIGHT
            ws.cell(row=r, column=5,
                    value=f"=F{r}*COS(RADIANS($B$3))+G{r}*SIN(RADIANS($B$3))")
            ws.cell(row=r, column=9, value=f"=IF(A{r}=\"\",{BIG},A{r})")
            ws.cell(row=r, column=10, value=f"=IF(D{r}=\"\",0,D{r})")
        else:
            blank = f'OR(A{r}="",A{p}="")'
            dl = (f'=IF({blank},"",DEGREES(ACOS(MIN(1,MAX(-1,'
                  f"COS(RADIANS(B{p}))*COS(RADIANS(B{r}))+"
                  f"SIN(RADIANS(B{p}))*SIN(RADIANS(B{r}))*COS(RADIANS(C{r}-C{p})))))))")
            ws.cell(row=r, column=11, value=dl)
            ws.cell(row=r, column=12,
                    value=f'=IF(K{r}="","",IF(K{r}=0,1,2/RADIANS(K{r})*TAN(RADIANS(K{r})/2)))')
            ws.cell(row=r, column=4,
                    value=f'=IF(K{r}="","",D{p}+(A{r}-A{p})/2*'
                          f"(COS(RADIANS(B{p}))+COS(RADIANS(B{r})))*L{r})")
            # ROUND(...,4) kills 1e-16 trig residue (COS(RADIANS(90))<>0) that
            # would otherwise auto-scale into a fake diagonal on the plan chart
            ws.cell(row=r, column=6,
                    value=f'=IF(K{r}="","",ROUND(F{p}+(A{r}-A{p})/2*'
                          f"(SIN(RADIANS(B{p}))*COS(RADIANS(C{p}))+SIN(RADIANS(B{r}))*COS(RADIANS(C{r})))*L{r},4))")
            ws.cell(row=r, column=7,
                    value=f'=IF(K{r}="","",ROUND(G{p}+(A{r}-A{p})/2*'
                          f"(SIN(RADIANS(B{p}))*SIN(RADIANS(C{p}))+SIN(RADIANS(B{r}))*SIN(RADIANS(C{r})))*L{r},4))")
            ws.cell(row=r, column=5,
                    value=f'=IF(D{r}="","",F{r}*COS(RADIANS($B$3))+G{r}*SIN(RADIANS($B$3)))')
            ws.cell(row=r, column=8,
                    value=f'=IF(K{r}="","",IF(A{r}=A{p},0,K{r}*100/(A{r}-A{p})))')
            ws.cell(row=r, column=9, value=f'=IF(A{r}="",{BIG},A{r})')
            ws.cell(row=r, column=10, value=f'=IF(D{r}="",0,D{r})')
        for col in range(4, 9):
            cell = ws.cell(row=r, column=col)
            cell.number_format = "#,##0.0" if col != 8 else "0.00"
            cell.alignment = S.AL_RIGHT
            if not (idx == DATA_FIRST and col in (4, 6, 7)):
                cell.font = S.F_LABEL
        # numeric chart helpers: NA() for unused rows so chart refs stay numeric
        for col, src in ((13, f"E{r}"), (14, f"D{r}"), (15, f"G{r}"), (16, f"F{r}")):
            ws.cell(row=r, column=col, value=f'=IF($A{r}="",NA(),{src})')
        for col in (9, 10, 11, 12, 13, 14, 15, 16):
            cell = ws.cell(row=r, column=col)
            cell.font = S.F_NOTE
            cell.number_format = "0.0"
            cell.alignment = S.AL_RIGHT

    note_row = DATA_LAST + 1
    ws.cell(row=note_row, column=1,
            value="* helper columns used by charts and the Kill Sheet - do not edit. "
                  "VS_c/TVD_c/EW_c/NS_c show #N/A on unused rows by design (chart gaps).").font = S.F_NOTE

    # ------------------------------------------------------------ interpolator
    r = DATA_LAST + 3
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    band = ws.cell(row=r, column=1, value="TVD Interpolator (same logic the Kill Sheet uses)")
    band.font = S.F_SECTION
    band.fill = S.FILL_SECTION
    for col in range(1, 9):
        ws.cell(row=r, column=col).fill = S.FILL_SECTION
    r += 1
    ws.cell(row=r, column=1, value="Measured depth").font = S.F_LABEL
    b = ws.cell(row=r, column=2, value=8000)
    b.fill = S.FILL_INPUT
    b.border = S.BORDER_BOX
    b.font = S.F_INPUT
    b.protection = UNLOCKED
    b.number_format = "#,##0"
    ws.cell(row=r, column=3, value="ft").font = S.F_UNIT
    md_ref = f"$B${r}"
    r += 1
    ws.cell(row=r, column=1, value="TVD at that MD").font = S.F_LABEL
    b = ws.cell(row=r, column=2, value="=" + interp_tvd(md_ref))
    b.fill = S.FILL_OUTPUT
    b.border = S.BORDER_BOX
    b.font = S.F_OUTPUT
    b.number_format = "#,##0.0"
    ws.cell(row=r, column=3, value="ft").font = S.F_UNIT

    # ------------------------------------------------------------ charts
    profile = ScatterChart()
    profile.varyColors = False  # else Excel colors every point and lists each in the legend
    profile.title = "Wellbore Profile (VS vs TVD)"
    profile.x_axis.title = "Vertical section (ft)"
    profile.y_axis.title = "TVD (ft)"
    profile.y_axis.scaling.orientation = "maxMin"
    for ax in (profile.x_axis, profile.y_axis):
        ax.delete = False
        ax.tickLblPos = "nextTo"
    profile.height = 10
    profile.width = 16
    xref = Reference(ws, min_col=13, min_row=DATA_FIRST, max_row=DATA_LAST)
    yref = Reference(ws, min_col=14, min_row=DATA_FIRST, max_row=DATA_LAST)
    s = Series(yref, xref, title="Trajectory")
    s.marker.symbol = "circle"
    s.marker.size = 4
    profile.series.append(s)
    ws.add_chart(profile, "R6")

    plan = ScatterChart()
    plan.varyColors = False
    plan.title = "Plan View (E/W vs N/S)"
    plan.x_axis.title = "E/W (ft)"
    plan.y_axis.title = "N/S (ft)"
    for ax in (plan.x_axis, plan.y_axis):
        ax.delete = False
        ax.tickLblPos = "nextTo"
    plan.height = 10
    plan.width = 16
    xref = Reference(ws, min_col=15, min_row=DATA_FIRST, max_row=DATA_LAST)
    yref = Reference(ws, min_col=16, min_row=DATA_FIRST, max_row=DATA_LAST)
    s = Series(yref, xref, title="Trajectory")
    s.marker.symbol = "circle"
    s.marker.size = 4
    plan.series.append(s)
    ws.add_chart(plan, "R27")
