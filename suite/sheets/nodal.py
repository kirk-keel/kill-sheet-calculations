"""Completions: nodal analysis - natural-flow operating point from IPR vs VLP.

Vogel IPR anchored on a single well test; simplified vertical-lift performance
(constant flowing gradient + rate-squared friction term). A 21-row rate table
computes both curves, the operating point is found by locating the first row
where the VLP curve rises above the IPR curve and linearly interpolating the
zero of (IPR - VLP) between that row and the previous one.

Table geometry:
  A=Q  B=IPR Pwf  C=VLP Pwf  D=diff (IPR-VLP)  E=flag (1 where diff<0)
IPR Pwf is NA()-guarded above Qmax so chart source ranges stay purely numeric.
"""
from openpyxl.chart import Reference, ScatterChart, Series

from .. import style as S
from ..builder import UNLOCKED, SheetBuilder

N_ROWS = 21


def _cell(ws, row, col, value, fmt="#,##0.00", is_input=False, is_output=False):
    c = ws.cell(row=row, column=col, value=value)
    c.number_format = fmt
    c.border = S.BORDER_BOX
    c.alignment = S.AL_RIGHT
    if is_input:
        c.fill = S.FILL_INPUT
        c.font = S.F_INPUT
        c.protection = UNLOCKED
    elif is_output:
        c.fill = S.FILL_OUTPUT
        c.font = S.F_OUTPUT
    else:
        c.font = S.F_LABEL
    return c


def build(wb):
    sb = SheetBuilder(
        wb, "Nodal Analysis", "completions",
        "Nodal Analysis - Natural Flow Operating Point",
        "Vogel IPR from a well test vs simplified tubing outflow (VLP); "
        "rate table, operating point and system plot",
        description="Vogel IPR vs simplified VLP, 21-point system curves, interpolated "
                    "natural-flow operating point and IPR/VLP chart",
    )
    ws = sb.ws

    sb.section("Reservoir Inflow (IPR)")

    c = sb.calc("Vogel IPR from Well Test",
                eq="Qmax = Qtest / (1 - 0.2 x (Pwf/Pr) - 0.8 x (Pwf/Pr)^2);   "
                   "Pwf(Q) = 0.125 x Pr x (SQRT(81 - 80 x Q/Qmax) - 1)",
                note="Vogel's dimensionless IPR for solution-gas-drive wells below the bubble "
                     "point. Anchor on a stabilized test rate and flowing BHP.")
    pr = c.link("Reservoir pressure (Pr)", "WD_ResP", "psi", fmt="#,##0")
    qt = c.inp("Test rate (Qtest)", 800, None, fmt="#,##0", note="STB/d, stabilized.")
    pwft = c.inp("Test flowing BHP (Pwf test)", 3300, "psi", fmt="#,##0")
    qmax = c.out("Vogel AOF (Qmax)",
                 f"IF({pwft}>={pr},NA(),"
                 f"{qt}/(1-0.2*({pwft}/{pr})-0.8*({pwft}/{pr})^2))",
                 None, fmt="#,##0",
                 note="STB/d at Pwf = 0. NA if the test point is invalid (Pwf >= Pr).")
    c.out("Test point validity",
          f'IF({pwft}>={pr},"CHECK: test Pwf >= reservoir pressure - '
          f'enter a valid stabilized test","OK - test point below Pr")',
          None, fmt="@")
    qtab = c.out("Rate-table maximum (Q axis span)",
                 f"IFERROR(ROUND(1.5*{qmax}/100,0)*100,NA())",
                 None, fmt="#,##0",
                 note="STB/d. Table below runs 0 to ~1.5 x Qmax in 20 even steps.")

    sb.section("Tubing Outflow (VLP - simplified)")

    c = sb.calc("Simplified Vertical Lift Performance",
                eq="Pwf_VLP = WHP + grad x TVD + C_f x (Q/1000)^2 x L/1000   (L = TVD here)",
                note="Average flowing gradient from gradient curves or a measured flowing "
                     "survey - this simplified VLP treats the hydrostatic term as "
                     "rate-independent, so it understates the lift-curve minimum at low "
                     "rates. Good for a first-pass operating point only.")
    whp = c.inp("Wellhead flowing pressure (WHP)", 250, "psi", fmt="#,##0")
    tvd = c.link("Perforation TVD (L)", "WD_PerfTVD", "ft", fmt="#,##0")
    tbgid = c.link("Tubing ID", "WD_TbgID", "in", fmt="0.000",
                   note="Linked from Well Data for reference - C_f below should reflect "
                        "this ID (smaller ID -> larger C_f).")
    grad = c.inp("Average flowing fluid gradient", 0.33, "psi/ft", fmt="0.000",
                 note="From gradient curves or measured flowing gradient survey.")
    cf = c.inp("Friction coefficient (C_f)", 45, None, fmt="0.0",
               note="psi per (MSTB/d)^2 per 1000 ft - documented rough default for "
                    "~3 in tubing; calibrate to a measured flowing point.")

    # ------------------------------------------------------------ rate table
    sb.row += 1
    band_row = sb.row
    ws.merge_cells(start_row=band_row, start_column=1, end_row=band_row, end_column=6)
    band = ws.cell(row=band_row, column=1,
                   value="System Curves - 21 rates from 0 to the table maximum "
                         "(all rows computed; edit the inputs above)")
    band.font = S.F_SECTION
    band.fill = S.FILL_SECTION
    for col in range(1, 7):
        ws.cell(row=band_row, column=col).fill = S.FILL_SECTION
    sb.row += 1

    hdr_row = sb.row
    headers = ["Q (STB/d)", "IPR Pwf (psi)", "VLP Pwf (psi)", "IPR - VLP (psi)", "flag*"]
    for i, h in enumerate(headers, start=1):
        hc = ws.cell(row=hdr_row, column=i, value=h)
        hc.font = S.F_HEADER
        hc.border = S.BORDER_BOX
        hc.alignment = S.AL_RIGHT if i > 1 else S.AL_LEFT
    sb.row += 1

    first = sb.row
    last = first + N_ROWS - 1
    for i in range(N_ROWS):
        r = first + i
        _cell(ws, r, 1, f"=IFERROR({i}/20*{qtab},NA())", fmt="#,##0")
        # IPR: Vogel inverse, NA() above Qmax / on invalid inputs so the
        # chart range stays numeric and errors degrade to gaps
        _cell(ws, r, 2,
              f"=IFERROR(IF($A{r}>{qmax},NA(),"
              f"0.125*{pr}*(SQRT(81-80*$A{r}/{qmax})-1)),NA())",
              fmt="#,##0")
        _cell(ws, r, 3,
              f"=IFERROR({whp}+{grad}*{tvd}+{cf}*($A{r}/1000)^2*{tvd}/1000,NA())",
              fmt="#,##0")
        _cell(ws, r, 4, f"=IF(ISNA(B{r}),NA(),B{r}-C{r})", fmt="#,##0")
        fc = _cell(ws, r, 5, f"=IF(ISNA(D{r}),0,IF(D{r}<0,1,0))", fmt="0")
        fc.font = S.F_NOTE
    sb.row = last + 1
    note = ws.cell(row=sb.row, column=1,
                   value="* flag = 1 where VLP exceeds IPR (used to locate the crossing). "
                         "IPR Pwf shows #N/A above Qmax by design (chart gap).")
    note.font = S.F_NOTE
    sb.row += 1

    q_rng = f"$A${first}:$A${last}"
    d_rng = f"$D${first}:$D${last}"
    flag_rng = f"$E${first}:$E${last}"

    sb.section("Operating Point")

    c = sb.calc("Natural-Flow Operating Point",
                eq="Qop: linear interpolation of (IPR - VLP) = 0 between the first row "
                   "with flag = 1 and the row before it",
                note="MATCH finds the first table row where the VLP curve is above the "
                     "IPR curve; the zero of the difference is interpolated between that "
                     "row and the previous one. If the curves never cross (or cross at "
                     "Q = 0) the well will not flow naturally.")
    m = c.out("First row with VLP above IPR", f"IFERROR(MATCH(1,{flag_rng},0),0)",
              None, fmt="0",
              note="Row index within the table; 0 = no crossing in the table range.")
    qop = c.out("Operating rate (Qop)",
                f"IF({m}<=1,NA(),INDEX({q_rng},{m}-1)+INDEX({d_rng},{m}-1)"
                f"*(INDEX({q_rng},{m})-INDEX({q_rng},{m}-1))"
                f"/(INDEX({d_rng},{m}-1)-INDEX({d_rng},{m})))",
                None, fmt="#,##0", note="STB/d. #N/A when there is no natural-flow point.")
    c.out("Operating flowing BHP (Pwf op)",
          f"IF({m}<=1,NA(),{whp}+{grad}*{tvd}+{cf}*({qop}/1000)^2*{tvd}/1000)",
          "psi", fmt="#,##0")
    c.out("Status",
          f'IF({m}<=1,"CHECK: no natural flow - artificial lift required",'
          f'"OK - natural flow operating point found")',
          None, fmt="@")

    # ------------------------------------------------------------ chart
    chart = ScatterChart()
    chart.varyColors = False
    chart.title = "IPR vs VLP - Operating Point"
    chart.x_axis.title = "Rate (STB/d)"
    chart.y_axis.title = "Flowing BHP (psi)"
    for ax in (chart.x_axis, chart.y_axis):
        ax.delete = False
        ax.tickLblPos = "nextTo"
    chart.height = 11
    chart.width = 18
    xref = Reference(ws, min_col=1, min_row=first, max_row=last)
    for col, name in ((2, "IPR (Vogel)"), (3, "VLP (simplified)")):
        yref = Reference(ws, min_col=col, min_row=first, max_row=last)
        s = Series(yref, xref, title=name)
        s.marker.symbol = "circle"
        s.marker.size = 4
        chart.series.append(s)
    ws.add_chart(chart, "H6")
