"""Shared layout helpers for the printable procedure-builder sheets.

Each procedure sheet is a print-ready step table whose description cells are
yellow (editable) Excel formulas that embed live workbook values via the
WD_* defined names, so the wording updates when Well Data changes and the
engineer can still overwrite any step.

Table layout:
  col A    step number
  col B:D  merged description (yellow formula cell, unlocked)
  col E    estimated hours (yellow, unlocked)
  col F    initials (blank, bordered)
"""
from openpyxl.styles import Font
from openpyxl.worksheet.properties import PageSetupProperties

from .. import style as S
from ..builder import UNLOCKED, SheetBuilder

COL_WIDTHS = {"A": 7, "B": 30, "C": 30, "D": 30, "E": 9, "F": 9}

TEMPLATE_NOTE = ("Template procedure - review and edit every step against the "
                 "program, permits and MOC before use.")

F_STEP = Font(name="Calibri", size=10)
F_WARN = Font(name="Calibri", size=9, bold=True, italic=True, color="9C5700")
F_SIGN = Font(name="Calibri", size=11)


class ProcedureSheet:
    """Thin wrapper over SheetBuilder for the procedure step tables."""

    def __init__(self, wb, name, group, title, subtitle="", description=""):
        self.sb = SheetBuilder(wb, name, group, title, subtitle, description)
        self.ws = self.sb.ws
        self.ws.sheet_view.showGridLines = False
        for col, w in COL_WIDTHS.items():
            self.ws.column_dimensions[col].width = w
        # repurpose the standard row-4 header for the step table
        headers = {1: "Step", 2: "Description (yellow = editable text)",
                   5: "Est. hrs", 6: "Initials"}
        for col in range(1, 7):
            c = self.ws.cell(row=4, column=col)
            c.value = headers.get(col, "")
            c.font = S.F_HEADER
            c.border = S.BORDER_BOX
            c.alignment = S.AL_CENTER if col != 2 else S.AL_LEFT
        self.ws.merge_cells("B4:D4")
        self.step_no = 0
        self._title_block()

    # ------------------------------------------------------------ title block
    def _pair(self, row, lcol, label, value, editable):
        ws = self.ws
        a = ws.cell(row=row, column=lcol, value=label)
        a.font = S.F_LABEL
        a.alignment = S.AL_LEFT
        ws.merge_cells(start_row=row, start_column=lcol + 1,
                       end_row=row, end_column=lcol + 2)
        b = ws.cell(row=row, column=lcol + 1, value=value)
        b.alignment = S.AL_LEFT
        b.number_format = "@"
        b.font = S.F_INPUT if editable else S.F_LINK
        fill = S.FILL_INPUT if editable else S.FILL_LINK
        for col in (lcol + 1, lcol + 2):
            c = ws.cell(row=row, column=col)
            c.border = S.BORDER_BOX
            c.fill = fill
            if editable:
                c.protection = UNLOCKED

    def _title_block(self):
        r = self.sb.row
        for i, (link_lbl, link_val, inp_lbl) in enumerate([
                ("Well", "=WD_Well", "Prepared by"),
                ("Field", "=WD_Field", "Date"),
                ("Rig", "=WD_Rig", "Revision")]):
            self._pair(r + i, 1, link_lbl, link_val, editable=False)
            self._pair(r + i, 4, inp_lbl, "", editable=True)
        r += 3
        self.ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        w = self.ws.cell(row=r, column=1, value=TEMPLATE_NOTE)
        w.font = F_WARN
        w.alignment = S.AL_WRAP
        self.sb.row = r + 1

    # ------------------------------------------------------------ table body
    def phase(self, text):
        """Section band between phases."""
        self.sb.section(text)

    def step(self, description, hrs=None, height=30):
        """One numbered step row. `description` may be a plain string or an
        Excel formula (leading '='). The cell is yellow and unlocked either
        way so the engineer can overwrite the generated wording."""
        self.step_no += 1
        ws = self.ws
        r = self.sb.row
        a = ws.cell(row=r, column=1, value=self.step_no)
        a.font = S.F_LABEL
        a.alignment = S.AL_CENTER
        a.border = S.BORDER_BOX
        a.number_format = "0"
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        d = ws.cell(row=r, column=2, value=description)
        d.font = F_STEP
        d.alignment = S.AL_WRAP
        for col in range(2, 5):
            c = ws.cell(row=r, column=col)
            c.fill = S.FILL_INPUT
            c.border = S.BORDER_BOX
            c.protection = UNLOCKED
        e = ws.cell(row=r, column=5, value=hrs)
        e.fill = S.FILL_INPUT
        e.font = S.F_INPUT
        e.border = S.BORDER_BOX
        e.alignment = S.AL_RIGHT
        e.number_format = "0.0"
        e.protection = UNLOCKED
        f = ws.cell(row=r, column=6)
        f.border = S.BORDER_BOX
        f.protection = UNLOCKED
        ws.row_dimensions[r].height = height
        self.sb.row += 1

    # ------------------------------------------------------------ close-out
    def finish(self):
        """Sign-off block, revision note row and print setup."""
        ws = self.ws
        r = self.sb.row + 1
        for label in ("Prepared by (engineer):", "Reviewed by (supervisor):",
                      "Approved by (operations superintendent):"):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
            c = ws.cell(row=r, column=1,
                        value=label + "  ____________________________")
            c.font = F_SIGN
            ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
            d = ws.cell(row=r, column=5, value="Date: ______________")
            d.font = F_SIGN
            ws.row_dimensions[r].height = 20
            r += 1
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        rev = ws.cell(row=r, column=1,
                      value="Rev 0 - template issue. Record revisions and MOC "
                            "references here before execution.")
        rev.font = S.F_INPUT
        rev.fill = S.FILL_INPUT
        rev.alignment = S.AL_WRAP
        for col in range(1, 7):
            c = ws.cell(row=r, column=col)
            c.fill = S.FILL_INPUT
            c.border = S.BORDER_BOX
            c.protection = UNLOCKED
        r += 1
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        n = ws.cell(row=r, column=1, value=TEMPLATE_NOTE)
        n.font = S.F_NOTE
        n.alignment = S.AL_WRAP
        self.sb.row = r + 1

        ws.print_area = f"A1:F{r}"
        ws.page_setup.orientation = "portrait"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
