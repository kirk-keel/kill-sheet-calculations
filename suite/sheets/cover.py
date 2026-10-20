"""Cover sheet: title, legend, and hyperlinked table of contents."""
from openpyxl.styles import Font, PatternFill, Alignment

from .. import style as S
from ..builder import SheetBuilder

GROUP_TITLES = {
    "reference": "Reference",
    "drilling": "Drilling",
    "completions": "Completions",
    "interventions": "Interventions",
}


def build(wb):
    ws = wb.create_sheet("Cover", 0)
    ws.sheet_properties.tabColor = S.TAB_COLORS["reference"]
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 34
    ws.column_dimensions["C"].width = 95

    ws.merge_cells("A1:C2")
    t = ws.cell(row=1, column=1, value="Drilling, Completion & Intervention Engineering Suite")
    t.font = Font(name="Calibri", size=22, bold=True, color="FFFFFF")
    t.fill = S.FILL_TITLE
    t.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    for r in (1, 2):
        for col in (1, 2, 3):
            ws.cell(row=r, column=col).fill = S.FILL_TITLE
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 24

    ws.cell(row=4, column=2, value="How to use").font = S.F_SECTION
    legend = [
        ("Yellow cells", "User inputs - overwrite the example values with your well data.", S.FILL_INPUT),
        ("Blue cells", "Linked to the Well Data master sheet - update the well there once, "
                       "or overwrite locally for a one-off value.", S.FILL_LINK),
        ("Green cells", "Live calculated results - do not type over them.", S.FILL_OUTPUT),
        ("Red / amber highlights", "Automatic alarms on status cells (limits exceeded, buckling, "
                                   "marginal margins).", None),
        ("SI columns", "Every value is mirrored to SI units automatically (columns D-E).", None),
        ("Notes / Equation", "The governing equation for every calculation is written out in column F.", None),
        ("Protection", "Sheets are protected with no password so formulas can't be typed over "
                       "by accident - Review > Unprotect Sheet to edit structure.", None),
    ]
    r = 5
    for label, text, fill in legend:
        b = ws.cell(row=r, column=2, value=label)
        b.font = S.F_LABEL
        if fill:
            b.fill = fill
            b.border = S.BORDER_BOX
        cc = ws.cell(row=r, column=3, value=text)
        cc.font = S.F_NOTE
        cc.alignment = S.AL_LEFT
        r += 1
    r += 1

    ws.cell(row=r, column=2,
            value="Field units are primary (ppg, psi, ft, in, bbl). All formulas are standard "
                  "API/IADC/IWCF-form oilfield equations; see each sheet's notes for assumptions.")
    ws.cell(row=r, column=2).font = S.F_NOTE
    r += 2

    by_group = {}
    for name, group, desc in SheetBuilder.registry:
        by_group.setdefault(group, []).append((name, desc))

    link_font = Font(name="Calibri", size=11, bold=True, color="0563C1", underline="single")
    for group in ("reference", "drilling", "completions", "interventions"):
        if group not in by_group:
            continue
        band = ws.cell(row=r, column=2, value=GROUP_TITLES[group])
        band.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
        band.fill = PatternFill("solid", fgColor=S.TAB_COLORS[group])
        ws.cell(row=r, column=3).fill = PatternFill("solid", fgColor=S.TAB_COLORS[group])
        r += 1
        for name, desc in by_group[group]:
            a = ws.cell(row=r, column=2, value=name)
            a.font = link_font
            a.hyperlink = f"#'{name}'!A1"
            d = ws.cell(row=r, column=3, value=desc)
            d.font = S.F_NOTE
            d.alignment = S.AL_LEFT
            r += 1
        r += 1

    ws.cell(row=r + 1, column=2,
            value="Engineering reference only - verify against current company standards and "
                  "manufacturer data before operational use.").font = S.F_NOTE
