"""Reference: API tubular catalog with lookup calculator.

Defines workbook names used elsewhere:
  TL_Names - catalog name column (dropdown lists)
  TL_Table - Name / OD / ID / weight table (VLOOKUP auto-fill)

Kill Sheet component names use these exact strings so their OD/ID auto-fill.
Dimensional data: API 5DP / 5CT nominal sizes; weights are nominal (drill pipe
adjusted weights incl. tool joints vary by connection - see note on sheet).
"""
from openpyxl.workbook.defined_name import DefinedName

from .. import style as S
from ..builder import SheetBuilder

SHEET = "Tubular Library"

# (name, OD in, ID in, nominal weight lb/ft)
CATALOG = [
    # drill pipe
    ("DP 3-1/2 13.30#", 3.500, 2.764, 13.30),
    ("DP 4 14.00#", 4.000, 3.340, 14.00),
    ("DP 4-1/2 16.60#", 4.500, 3.826, 16.60),
    ("DP 5 19.50#", 5.000, 4.276, 19.50),
    ("DP 5 25.60#", 5.000, 4.000, 25.60),
    ("DP 5-1/2 21.90#", 5.500, 4.778, 21.90),
    ("DP 5-1/2 24.70#", 5.500, 4.670, 24.70),
    ("DP 6-5/8 25.20#", 6.625, 5.965, 25.20),
    # heavy weight drill pipe
    ("HWDP 3-1/2 25.3#", 3.500, 2.063, 25.30),
    ("HWDP 4 29.7#", 4.000, 2.563, 29.70),
    ("HWDP 4-1/2 41.0#", 4.500, 2.750, 41.00),
    ("HWDP 5 49.3#", 5.000, 3.000, 49.30),
    ("HWDP 5-1/2 57.0#", 5.500, 3.375, 57.00),
    # drill collars
    ("DC 4-3/4 x 2-1/4", 4.750, 2.250, 46.70),
    ("DC 6-1/4 x 2-13/16", 6.250, 2.813, 82.50),
    ("DC 6-1/2 92#", 6.500, 2.813, 91.60),
    ("DC 8 x 2-13/16", 8.000, 2.813, 150.50),
    ("DC 9-1/2 x 3", 9.500, 3.000, 217.00),
    # casing
    ("CSG 13-3/8 68#", 13.375, 12.415, 68.00),
    ("CSG 10-3/4 51#", 10.750, 9.850, 51.00),
    ("CSG 9-5/8 43.5#", 9.625, 8.755, 43.50),
    ("CSG 9-5/8 47#", 9.625, 8.681, 47.00),
    ("CSG 9-5/8 53.5#", 9.625, 8.535, 53.50),
    ("CSG 7-5/8 39#", 7.625, 6.625, 39.00),
    ("CSG 7 26#", 7.000, 6.276, 26.00),
    ("CSG 7 29#", 7.000, 6.184, 29.00),
    ("CSG 7 32#", 7.000, 6.094, 32.00),
    ("CSG 5-1/2 17#", 5.500, 4.892, 17.00),
    ("CSG 5-1/2 20#", 5.500, 4.778, 20.00),
    ("CSG 5-1/2 23#", 5.500, 4.670, 23.00),
    # tubing
    ("TBG 2-3/8 4.7#", 2.375, 1.995, 4.70),
    ("TBG 2-7/8 6.5#", 2.875, 2.441, 6.50),
    ("TBG 3-1/2 9.3#", 3.500, 2.992, 9.30),
    ("TBG 4 11#", 4.000, 3.476, 11.00),
    ("TBG 4-1/2 12.75#", 4.500, 3.958, 12.75),
    # coiled tubing (wall from common gauges)
    ("CT 1-1/2 x 0.109w", 1.500, 1.282, 1.62),
    ("CT 1-3/4 x 0.134w", 1.750, 1.482, 2.31),
    ("CT 2 x 0.156w", 2.000, 1.688, 3.07),
    ("CT 2-3/8 x 0.175w", 2.375, 2.025, 4.11),
]

DATA_FIRST = 9  # banner 1-2, col headers 4, section band 6, note 7, table header 8


def build(wb):
    sb = SheetBuilder(
        wb, SHEET, "reference",
        "Tubular Library - API Catalog",
        "Standard drill pipe, HWDP, collars, casing, tubing and CT dimensions. "
        "Feeds the dropdowns and OD/ID auto-fill on the Kill Sheet.",
        description="API tubular catalog (TL_Names/TL_Table named ranges), capacities/"
                    "displacements per size, and a lookup calculator",
    )
    ws = sb.ws
    ws.freeze_panes = f"A{DATA_FIRST}"

    sb.section("Catalog")
    ws.cell(row=sb.row, column=1,
            value="Nominal API dimensions. DP weights are nominal - for string weight use "
                  "adjusted weight (incl. tool joints, ~7% higher). Pressure ratings are "
                  "grade-specific: enter them on the Casing & Tubular Design sheet.").font = S.F_NOTE
    sb.row += 1
    hdr = sb.row
    headers = ["Name", "OD (in)", "ID (in)", "Weight (lb/ft)", "Capacity (bbl/ft)",
               "Disp. closed (bbl/ft)", "Disp. open (bbl/ft)", "Wall (in)"]
    for i, h in enumerate(headers, start=1):
        cell = ws.cell(row=hdr, column=i, value=h)
        cell.font = S.F_HEADER
        cell.border = S.BORDER_BOX
    first = hdr + 1
    assert first == DATA_FIRST, (first, DATA_FIRST)
    for i, (name, od, idd, wt) in enumerate(CATALOG):
        r = first + i
        ws.cell(row=r, column=1, value=name).border = S.BORDER_BOX
        for col, v, fmt in ((2, od, "0.000"), (3, idd, "0.000"), (4, wt, "0.00")):
            cell = ws.cell(row=r, column=col, value=v)
            cell.number_format = fmt
            cell.border = S.BORDER_BOX
        for col, f, in ((5, f"=$C{r}^2/1029.4"), (6, f"=$B{r}^2/1029.4"),
                        (7, f"=($B{r}^2-$C{r}^2)/1029.4"), (8, f"=($B{r}-$C{r})/2")):
            cell = ws.cell(row=r, column=col, value=f)
            cell.number_format = "0.00000" if col < 8 else "0.000"
            cell.border = S.BORDER_BOX
            cell.font = S.F_SI
    last = first + len(CATALOG) - 1
    sb.row = last + 1

    wb.defined_names["TL_Names"] = DefinedName(
        "TL_Names", attr_text=f"'{SHEET}'!$A${first}:$A${last}")
    wb.defined_names["TL_Table"] = DefinedName(
        "TL_Table", attr_text=f"'{SHEET}'!$A${first}:$D${last}")

    sb.section("Lookup Calculator")
    c = sb.calc("Pipe Properties Lookup",
                eq="Capacity = ID^2/1029.4;  Disp closed = OD^2/1029.4;  "
                   "Disp open = (OD^2-ID^2)/1029.4",
                note="Type or pick any catalog name (dropdowns available on the Kill Sheet).")
    nm = c.inp("Catalog name", "DP 5 19.50#", None, fmt="@")
    od = c.out("OD", f"VLOOKUP({nm},TL_Table,2,FALSE)", "in", fmt="0.000")
    idd = c.out("ID", f"VLOOKUP({nm},TL_Table,3,FALSE)", "in", fmt="0.000")
    c.out("Nominal weight", f"VLOOKUP({nm},TL_Table,4,FALSE)", "lb/ft", fmt="0.00")
    c.out("Capacity", f"{idd}^2/1029.4", "bbl/ft", fmt="0.00000")
    c.out("Displacement (closed end)", f"{od}^2/1029.4", "bbl/ft", fmt="0.00000")
    c.out("Displacement (open end)", f"({od}^2-{idd}^2)/1029.4", "bbl/ft", fmt="0.00000")
    hole = c.inp("Hole / casing ID for annulus", 8.681, "in", fmt="0.000")
    c.out("Annular capacity around pipe", f"({hole}^2-{od}^2)/1029.4", "bbl/ft", fmt="0.00000")
