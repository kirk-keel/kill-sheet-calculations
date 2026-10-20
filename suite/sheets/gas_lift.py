"""Completions: continuous gas lift design - valve spacing, gas rate, dome pressures.

Classic load-fluid (kill-fluid) graphical spacing implemented as a 10-row
table: valve 1 goes where kick-off casing pressure balances the static load
fluid column; each deeper valve is spaced by the gap between the gas-charged
casing pressure and the unloading tubing pressure at the valve above, less a
transfer margin. Casing surface pressure is dropped 'drop-per-valve' for each
prior valve. Depths cap at perf TVD - 100 ft; rows past the cap show #N/A so
the spacing-diagram chart ranges stay purely numeric.

Table geometry:
  A=valve #  B=depth  C=Pcsg surface for valve  D=Pcsg at depth
  E=Pt (unloading tubing) at depth  F=spacing to next  G=depth helper (NA->0)
"""
from openpyxl.chart import Reference, ScatterChart, Series

from .. import style as S
from ..builder import SheetBuilder
from .nodal import _cell

N_VALVES = 10


def build(wb):
    sb = SheetBuilder(
        wb, "Gas Lift Design", "completions",
        "Continuous Gas Lift Design",
        "Load-fluid valve spacing with gas-column casing pressure, injection "
        "gas requirement and nitrogen dome pressure correction",
        description="Gas lift: 10-valve load-fluid spacing table with spacing "
                    "diagram, injection gas rate and valve test-rack pressures",
    )
    ws = sb.ws

    sb.section("Design Inputs")

    c = sb.calc("Well & Design Data",
                eq="Pwf = Pres - Q/J",
                note="Tubing/casing IDs are linked for reference - the "
                     "unloading gradient and target GLR below come from "
                     "gradient curves for this tubing size and rate.")
    pres = c.link("Reservoir pressure (Pres)", "WD_ResP", "psi", fmt="#,##0")
    perf = c.link("Perforation TVD", "WD_PerfTVD", "ft", fmt="#,##0")
    c.link("Tubing ID", "WD_TbgID", "in", fmt="0.000")
    c.link("Casing ID", "WD_CsgID", "in", fmt="0.000")
    qt = c.inp("Target rate (Q)", 1200, None, fmt="#,##0", note="bpd liquid.")
    j = c.inp("Productivity index (J)", 1.0, None, fmt="0.00", note="bpd/psi.")
    pwh = c.inp("Wellhead flowing pressure (Pwh)", 100, "psi", fmt="#,##0")
    pko = c.inp("Kick-off / injection surface pressure (Pko)", 1000, "psi",
                fmt="#,##0", note="Available casing-head injection pressure.")
    sgg = c.inp("Injection gas specific gravity", 0.65, None, fmt="0.00",
                note="Air = 1.0.")
    tsurf = c.inp("Average surface temperature", 80, None, fmt="0.0",
                  note="degF.", si=("degC", lambda ref: f"=({ref}-32)/1.8"))
    tgrad = c.inp("Temperature gradient", 1.6, None, fmt="0.00",
                  note="degF per 100 ft.")
    gs = c.inp("Load (kill) fluid gradient (Gs)", 0.465, "psi/ft", fmt="0.000",
               note="Static gradient of the fluid the well is loaded with.")
    gf = c.inp("Unloading tubing gradient (Gf)", 0.15, "psi/ft", fmt="0.000",
               note="Flowing gradient above the injection point - from "
                    "gradient curves for the target rate and tubing size.")
    drop = c.inp("Casing pressure drop per valve", 25, "psi", fmt="#,##0",
                 note="Surface injection pressure is reduced this much for "
                      "each valve already uncovered.")
    marg = c.inp("Valve transfer dP margin", 50, "psi", fmt="#,##0",
                 note="Casing-to-tubing differential reserved at each transfer.")
    c.out("Flowing BHP at target rate (Pwf)", f"{pres}-{qt}/{j}",
          "psi", fmt="#,##0")
    c.txt("Target injection depth",
          "As deep as the casing pressure allows - the deepest valve from the "
          "spacing table below is the operating (injection) point.")

    sb.section("Gas Column & Casing Pressure")

    c = sb.calc("Static Gas Column - Worked Example",
                eq="Pcsg(D) = Psurf x EXP(0.01875 x SG x D / (Tavg_R x Z));   "
                   "Tavg_R = 460 + Tsurf + Tgrad x D/200",
                note="Weight of the static injection-gas column adds to the "
                     "surface casing pressure. Tavg_R is the mean of surface "
                     "and at-depth temperatures in Rankine (simplified "
                     "average-column temperature). Z is an average gas "
                     "deviation factor for the column.")
    z = c.inp("Gas deviation factor (Z, average)", 0.9, None, fmt="0.00")
    dex = c.inp("Example depth", 6000, "ft", fmt="#,##0")
    tavg = c.out("Average column temperature (Tavg_R)",
                 f"460+{tsurf}+{tgrad}*{dex}/200", None, fmt="0.0",
                 note="degR.")
    c.out("Casing pressure at example depth",
          f"{pko}*EXP(0.01875*{sgg}*{dex}/({tavg}*{z}))", "psi", fmt="#,##0",
          note="At Pko surface pressure - the table below repeats this per "
               "valve with the reduced surface pressure.")

    # ------------------------------------------------- valve spacing table
    sb.row += 1
    band_row = sb.row
    ws.merge_cells(start_row=band_row, start_column=1, end_row=band_row,
                   end_column=6)
    band = ws.cell(row=band_row, column=1,
                   value="Valve Spacing Table - load-fluid design "
                         "(D1 = (Pko - Pwh)/Gs; Dn = Dn-1 + "
                         "(Pcsg@Dn-1 - Pt@Dn-1 - margin)/Gs)")
    band.font = S.F_SECTION
    band.fill = S.FILL_SECTION
    for col in range(1, 7):
        ws.cell(row=band_row, column=col).fill = S.FILL_SECTION
    sb.row += 1

    hdr_row = sb.row
    headers = ["Valve #", "Depth (ft)", "Pcsg surf (psi)", "Pcsg @ depth (psi)",
               "Pt @ depth (psi)", "Spacing to next (ft)", "depth*"]
    for i, h in enumerate(headers, start=1):
        hc = ws.cell(row=hdr_row, column=i, value=h)
        hc.font = S.F_HEADER
        hc.border = S.BORDER_BOX
        hc.alignment = S.AL_RIGHT if i > 1 else S.AL_LEFT
    sb.row += 1

    def pcsg(sref, dref):
        return (f"{sref}*EXP(0.01875*{sgg}*{dref}"
                f"/((460+{tsurf}+{tgrad}*{dref}/200)*{z}))")

    first = sb.row
    last = first + N_VALVES - 1
    for i in range(N_VALVES):
        r = first + i
        n = i + 1
        if n == 1:
            _cell(ws, r, 1, 1, fmt="0")
            _cell(ws, r, 2, f"=MIN(({pko}-{pwh})/{gs},{perf}-100)", fmt="#,##0.0")
        else:
            p = r - 1
            _cell(ws, r, 1, f"=IF(ISNA($B{r}),NA(),{n})", fmt="0")
            _cell(ws, r, 2,
                  f"=IF(ISNA($B{p}),NA(),IF(OR($B{p}>={perf}-100,$F{p}<=0),"
                  f"NA(),MIN($B{p}+$F{p},{perf}-100)))",
                  fmt="#,##0.0")
        _cell(ws, r, 3, f"=IF(ISNA($B{r}),NA(),{pko}-{n - 1}*{drop})",
              fmt="#,##0")
        _cell(ws, r, 4,
              f"=IF(ISNA($B{r}),NA()," + pcsg(f"$C{r}", f"$B{r}") + ")",
              fmt="#,##0")
        _cell(ws, r, 5, f"=IF(ISNA($B{r}),NA(),{pwh}+{gf}*$B{r})", fmt="#,##0")
        _cell(ws, r, 6,
              f"=IF(ISNA($B{r}),NA(),($D{r}-$E{r}-{marg})/{gs})", fmt="#,##0.0")
        hcell = _cell(ws, r, 7, f"=IF(ISNA($B{r}),0,$B{r})", fmt="#,##0.0")
        hcell.font = S.F_NOTE
    sb.row = last + 1
    note = ws.cell(row=sb.row, column=1,
                   value="* helper column: depth with #N/A rows as 0, so MAX() "
                         "below returns the deepest valve. Depths cap at perf "
                         "TVD - 100 ft; rows past the cap show #N/A by design "
                         "(chart gap). Pcsg surf drops 'drop-per-valve' for "
                         "each prior valve.")
    note.font = S.F_NOTE
    sb.row += 1

    c = sb.calc("Operating (Deepest) Valve")
    opd = c.out("Operating valve depth",
                f"MAX($G${first}:$G${last})", "ft", fmt="#,##0.0",
                note="Deepest numeric depth in the table - the design "
                     "injection point.")
    c.out("Depth status",
          f'IF({opd}>=0.85*({perf}-100),"OK - injection point near target '
          f'depth",IF({opd}>=0.6*({perf}-100),"MARGINAL - injection point '
          f'well above target depth","CHECK: INSUFFICIENT casing pressure - '
          f'injection point shallow"))',
          None, fmt="@")

    # ------------------------------------------------------------ chart
    chart = ScatterChart()
    chart.varyColors = False
    chart.title = "Valve Spacing Diagram"
    chart.x_axis.title = "Pressure (psi)"
    chart.y_axis.title = "Depth (ft)"
    for ax in (chart.x_axis, chart.y_axis):
        ax.delete = False
        ax.tickLblPos = "nextTo"
    chart.y_axis.scaling.orientation = "maxMin"
    chart.height = 11
    chart.width = 16
    yref = Reference(ws, min_col=2, min_row=first, max_row=last)
    for col, name in ((4, "Pcsg at valve"), (5, "Pt at valve")):
        xref = Reference(ws, min_col=col, min_row=first, max_row=last)
        s = Series(yref, xref, title=name)
        s.marker.symbol = "circle"
        s.marker.size = 5
        chart.series.append(s)
    ws.add_chart(chart, "H6")

    sb.section("Injection Gas Requirement")

    c = sb.calc("Injection Gas Rate",
                eq="inj GLR = target GLR - formation GLR;   "
                   "Qgi [Mscf/d] = inj GLR x Q / 1000",
                note="Target total GLR from gradient curves for the target "
                     "rate and tubing size - the GLR that gives the flattest "
                     "practical flowing gradient.")
    tglr = c.inp("Target total GLR", 400, "scf/bbl", fmt="#,##0",
                 note="From gradient curves.")
    fglr = c.inp("Formation GLR", 120, "scf/bbl", fmt="#,##0",
                 note="Gas the reservoir already produces per barrel.")
    iglr = c.out("Injection GLR required", f"{tglr}-{fglr}", "scf/bbl",
                 fmt="#,##0")
    c.out("Injection gas rate", f"{iglr}*{qt}/1000", None, fmt="#,##0.0",
          note="Mscf/d.")
    pcd = c.inp("Compressor discharge pressure available", 1200, "psi",
                fmt="#,##0")
    c.out("Compressor status",
          f'IF({pcd}>={pko}+100,"OK - discharge pressure adequate for '
          f'kick-off",IF({pcd}>={pko},"MARGINAL - little margin over '
          f'kick-off pressure","CHECK: INSUFFICIENT compressor discharge '
          f'vs kick-off pressure"))',
          None, fmt="@")

    sb.section("Valve Dome Pressure Correction")

    c = sb.calc("Nitrogen Dome Pressure & Test Rack Setting",
                eq="Pd@depth = Pcsg@valve - margin;   "
                   "Ptro = Pd@depth x 520 / T_R@depth",
                note="Nitrogen-charged bellows valve: dome pressure set so the "
                     "valve holds casing pressure less the transfer margin at "
                     "the operating depth. Temperature correction to the 60 "
                     "degF test rack is the simplified ideal-gas ratio, with "
                     "port-to-bellows area ratio R = 0 assumed (unbalanced "
                     "effects ignored).")
    tv = c.out("Valve temperature at operating depth",
               f"{tsurf}+{tgrad}*{opd}/100", None, fmt="0.0",
               note="degF. T = Tsurf + Tgrad x D/100 at the operating valve.",
               si=("degC", lambda ref: f"=({ref}-32)/1.8"))
    pcv = c.out("Casing pressure at operating valve",
                pcsg(pko, opd), "psi", fmt="#,##0",
                note="Gas column at Pko surface pressure (operating surface "
                     "pressure assumed = Pko; simplification).")
    pd = c.out("Dome pressure required at depth",
               f"{pcv}-{marg}", "psi", fmt="#,##0")
    c.out("Test rack opening pressure at 60 degF (Ptro)",
          f"{pd}*520/(460+{tv})", "psi", fmt="#,##0",
          note="520 degR = 60 degF reference.")
