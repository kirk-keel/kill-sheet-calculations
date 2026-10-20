"""Reference: wellbore schematic creator.

Enter casing strings (OD, top, shoe), tubing, packer, TRSV, perfs and PBTD;
the sheet draws a scaled schematic on an XY scatter chart: depth (MD) on the
inverted Y axis, radial position (inches from centerline) on X. Casing walls
are line pairs, shoes are triangles, the packer is a square pair, perfs are
tick marks. Deviated wells draw "unwrapped" along MD.
"""
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties

from .. import style as S
from ..builder import UNLOCKED, SheetBuilder

N_CSG = 6

CSG_DEFAULTS = [
    ("Conductor 20in", 20.0, 94.0, 0, 120),
    ("Surface 13-3/8in", 13.375, 68.0, 0, 1500),
    ("Intermediate 9-5/8in 47#", 9.625, 47.0, 0, 5900),
    ("Production liner 5-1/2in 17#", 5.5, 17.0, 5600, 10450),
]

N_PERF_TICKS = 8


def _hdr(ws, row, headers, start_col=1):
    for i, h in enumerate(headers, start=start_col):
        c = ws.cell(row=row, column=i, value=h)
        c.font = S.F_HEADER
        c.border = S.BORDER_BOX
        c.alignment = S.AL_RIGHT if i > start_col else S.AL_LEFT


def _inp(ws, row, col, value, fmt="#,##0"):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = S.FILL_INPUT
    c.border = S.BORDER_BOX
    c.font = S.F_INPUT
    c.protection = UNLOCKED
    c.number_format = fmt
    c.alignment = S.AL_RIGHT
    return c


def build(wb):
    sb = SheetBuilder(
        wb, "Wellbore Schematic", "reference",
        "Wellbore Schematic Creator",
        "Enter casing, tubing, packer and perf depths (yellow) - the schematic chart "
        "draws itself to scale. Depths in MD; deviated wells draw unwrapped.",
        description="Customizable wellbore schematic: casing strings, tubing, packer, "
                    "TRSV, perfs, PBTD drawn to scale on a chart",
    )
    ws = sb.ws

    # ------------------------------------------------------------ casing table
    sb.section("Casing Strings (up to 6; leave unused rows blank)")
    hdr = sb.row
    _hdr(ws, hdr, ["String", "OD (in)", "Weight (lb/ft)", "Top MD (ft)", "Shoe MD (ft)"])
    cfirst = hdr + 1
    for i in range(N_CSG):
        r = cfirst + i
        name, od, wt, top, shoe = CSG_DEFAULTS[i] if i < len(CSG_DEFAULTS) else ("", None, None, None, None)
        nc = ws.cell(row=r, column=1, value=name)
        nc.fill = S.FILL_INPUT
        nc.border = S.BORDER_BOX
        nc.font = S.F_INPUT
        nc.protection = UNLOCKED
        _inp(ws, r, 2, od, "0.000")
        _inp(ws, r, 3, wt, "0.00")
        _inp(ws, r, 4, top, "#,##0")
        _inp(ws, r, 5, shoe, "#,##0")
    clast = cfirst + N_CSG - 1
    sb.row = clast + 1

    # ------------------------------------------------------------ completion
    sb.section("Completion & Wellbore Items")
    c = sb.calc("Tubing")
    tod = c.inp("Tubing OD", 3.5, "in", fmt="0.000")
    ttop = c.inp("Tubing top MD", 0, "ft", fmt="#,##0")
    tbot = c.inp("Tubing bottom MD (packer/seal depth)", 5480, "ft", fmt="#,##0")
    c = sb.calc("Packer & TRSV")
    pod = c.inp("Packer OD (gauge)", 8.0, "in", fmt="0.000")
    pdep = c.inp("Packer depth MD", 5480, "ft", fmt="#,##0")
    vdep = c.inp("TRSV depth MD", 1500, "ft", fmt="#,##0")
    c = sb.calc("Perforations & PBTD")
    ptop = c.inp("Perf top MD", 5950, "ft", fmt="#,##0")
    pbot = c.inp("Perf bottom MD", 10400, "ft", fmt="#,##0")
    pcod = c.inp("Perf interval casing OD", 5.5, "in", fmt="0.000")
    pbtd = c.inp("PBTD / TD MD", 10450, "ft", fmt="#,##0")

    c = sb.calc("Summary")
    c.out("Deepest shoe", f"MAX($E${cfirst}:$E${clast})", "ft", fmt="#,##0")
    c.out("Strings entered", f"COUNT($B${cfirst}:$B${clast})", None, fmt="0")
    c.txt("", "Chart at right redraws automatically. Depths are MD - a horizontal lateral "
              "shows as continued depth; annotate the KOP on the printout if needed.")

    # ------------------------------------------------------------ chart helpers
    hrow = 5  # helper block starts high in far columns
    for col, w in (("H", 4), ("I", 9), ("J", 9), ("K", 9), ("L", 9), ("M", 9),
                   ("N", 9), ("O", 9), ("P", 9), ("Q", 9), ("R", 9)):
        ws.column_dimensions[col].width = w
    ws.cell(row=hrow - 1, column=8, value="chart helpers - do not edit").font = S.F_NOTE

    # casing walls: 6 helper rows per string (left pair, gap, right pair, gap)
    cw_first = hrow
    for i in range(N_CSG):
        src = cfirst + i
        base = cw_first + i * 6
        pat = [(f"-$B${src}/2", f"$D${src}"), (f"-$B${src}/2", f"$E${src}"),
               (None, None),
               (f"$B${src}/2", f"$D${src}"), (f"$B${src}/2", f"$E${src}"),
               (None, None)]
        for j, (x, y) in enumerate(pat):
            r = base + j
            if x is None:
                ws.cell(row=r, column=9, value="=NA()")
                ws.cell(row=r, column=10, value="=NA()")
            else:
                ws.cell(row=r, column=9, value=f'=IF($B${src}="",NA(),{x})')
                ws.cell(row=r, column=10, value=f'=IF($B${src}="",NA(),{y})')
    cw_last = cw_first + N_CSG * 6 - 1

    # casing shoes: 2 marker rows per string
    sh_first = cw_last + 1
    for i in range(N_CSG):
        src = cfirst + i
        for j, sign in enumerate(("-", "")):
            r = sh_first + i * 2 + j
            ws.cell(row=r, column=9, value=f'=IF($B${src}="",NA(),{sign}$B${src}/2)')
            ws.cell(row=r, column=10, value=f'=IF($B${src}="",NA(),$E${src})')
    sh_last = sh_first + N_CSG * 2 - 1

    # tubing walls: 6 rows
    tw_first = sh_last + 1
    pat = [(f"-{tod}/2", ttop), (f"-{tod}/2", tbot), (None, None),
           (f"{tod}/2", ttop), (f"{tod}/2", tbot), (None, None)]
    for j, (x, y) in enumerate(pat):
        r = tw_first + j
        ws.cell(row=r, column=9, value="=NA()" if x is None else f"={x}")
        ws.cell(row=r, column=10, value="=NA()" if y is None else f"={y}")
    tw_last = tw_first + 5

    # packer: 2 square markers; TRSV: 2 diamonds; PBTD: 1 marker at centerline
    eq_first = tw_last + 1
    eq_rows = [(f"=-{pod}/2", f"={pdep}"), (f"={pod}/2", f"={pdep}"),
               (f"=-{tod}/2-0.6", f"={vdep}"), (f"={tod}/2+0.6", f"={vdep}"),
               ("=0", f"={pbtd}")]
    for j, (x, y) in enumerate(eq_rows):
        ws.cell(row=eq_first + j, column=9, value=x)
        ws.cell(row=eq_first + j, column=10, value=y)
    pk_first, pk_last = eq_first, eq_first + 1
    vv_first, vv_last = eq_first + 2, eq_first + 3
    bd_row = eq_first + 4

    # perforation ticks: 8 depths x both sides
    pf_first = bd_row + 1
    for j in range(N_PERF_TICKS * 2):
        r = pf_first + j
        k = j % N_PERF_TICKS
        sign = "-" if j >= N_PERF_TICKS else ""
        ws.cell(row=r, column=9, value=f"={sign}{pcod}/2")
        ws.cell(row=r, column=10, value=f"={ptop}+({pbot}-{ptop})*{k}/{N_PERF_TICKS - 1}")
    pf_last = pf_first + N_PERF_TICKS * 2 - 1

    # ------------------------------------------------------------ the chart
    chart = ScatterChart()
    chart.varyColors = False
    chart.title = "Wellbore Schematic"
    chart.x_axis.title = "Inches from centerline"
    chart.y_axis.title = "MD (ft)"
    chart.y_axis.scaling.orientation = "maxMin"
    for ax in (chart.x_axis, chart.y_axis):
        ax.delete = False
        ax.tickLblPos = "nextTo"
    chart.height = 22
    chart.width = 13

    # NA() points are INTERPOLATED (not gapped) in Excel scatter lines, so each
    # wall is its own 2-point series; marker series get line.noFill.
    def wall(first, color):
        xref = Reference(ws, min_col=9, min_row=first, max_row=first + 1)
        yref = Reference(ws, min_col=10, min_row=first, max_row=first + 1)
        s = Series(yref, xref)
        s.graphicalProperties.line = LineProperties(solidFill=color, w=28575)  # 2.25pt
        s.marker = Marker(symbol="none")
        return s

    def markers(first, last, name, sym, size):
        xref = Reference(ws, min_col=9, min_row=first, max_row=last)
        yref = Reference(ws, min_col=10, min_row=first, max_row=last)
        s = Series(yref, xref, title=name)
        s.marker = Marker(symbol=sym, size=size)
        s.graphicalProperties.line = LineProperties(noFill=True)
        return s

    for i in range(N_CSG):
        base = cw_first + i * 6
        chart.series.append(wall(base, "44546A"))       # left wall
        chart.series.append(wall(base + 3, "44546A"))   # right wall
    chart.series.append(wall(tw_first, "70AD47"))       # tubing left
    chart.series.append(wall(tw_first + 3, "70AD47"))   # tubing right
    chart.series.append(markers(sh_first, sh_last, "Shoes", "triangle", 8))
    chart.series.append(markers(pk_first, pk_last, "Packer", "square", 11))
    chart.series.append(markers(vv_first, vv_last, "TRSV", "diamond", 8))
    chart.series.append(markers(bd_row, bd_row, "PBTD", "x", 9))
    chart.series.append(markers(pf_first, pf_last, "Perfs", "dash", 6))
    chart.legend = None
    ws.add_chart(chart, "S4")
    ws.cell(row=clast + 2, column=8,
            value="Chart symbols: dark lines = casing, green = tubing, triangles = shoes, "
                  "square pair = packer, diamonds = TRSV, dashes = perfs, x = PBTD."
            ).font = S.F_NOTE
