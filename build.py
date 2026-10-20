"""Build the petroleum engineering suite workbook.

Usage:  python build.py  ->  petroleum_engineering_suite.xlsx
"""
import importlib

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

# Sheet modules in tab order. The cover is built last (it needs the registry)
# but inserted at position 0.
SHEET_MODULES = [
    # reference
    "well_data",
    "tubular_library",
    "units",
    # drilling
    "mud_hydrostatics",
    "capacities",
    "kill_sheet",
    "kill_form",
    "subsea_kill",
    "well_control",
    "hydraulics",
    "bit_hydraulics",
    "drill_string",
    "directional",
    "directional_survey",
    "torque_drag",
    "hole_cleaning",
    "trip_schedule",
    "pressure_window",
    "log_interpretation",
    "cementing",
    "casing_design",
    "procedure_drilling",
    # completions
    "completion_fluids",
    "tubing_packers",
    "tubing_stress",
    "perforating",
    "stimulation",
    "inflow",
    "nodal",
    "esp_design",
    "gas_lift",
    "rod_pump",
    "completion_design",
    "prod_chemistry",
    "completion_equipment",
    "advanced_completions",
    "procedure_completion",
    # interventions
    "rig_up",
    "coiled_tubing",
    "wireline",
    "snubbing",
    "well_kill_ops",
    "fishing",
    "procedure_intervention",
    # reference (drawn last so completion depths exist for defaults)
    "wellbore_schematic",
]

OUT = "petroleum_engineering_suite.xlsx"

# status-text keywords colored by conditional formatting on every sheet
BAD_WORDS = ["EXCEED", "CHECK", "DO NOT", "ABOVE FRAC", "ADD STEM", "BUCKLING",
             "INSUFFICIENT", "BELOW MINIMUM", "<> bit"]
RED_FILL = PatternFill("solid", fgColor="F8CBAD")
RED_FONT = Font(color="9C0006", bold=True)
AMBER_FILL = PatternFill("solid", fgColor="FFE699")
AMBER_FONT = Font(color="9C5700", bold=True)
GREEN_FONT = Font(color="1E6B2E", bold=True)


def add_status_formatting(ws):
    """Color status cells: red = alarm words, amber = MARGINAL, green = OK.
    Applied to the value column (B) and the T&D buckling-flag column (M) only,
    so instructional notes ('do not edit...') don't trigger alarms."""
    for rng, anchor in (("B1:B400", "B1"), ("M1:M400", "M1")):
        bad = "OR(" + ",".join(f'ISNUMBER(SEARCH("{w}",{anchor}))' for w in BAD_WORDS) + ")"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f"AND(ISTEXT({anchor}),{bad})"], fill=RED_FILL, font=RED_FONT))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND(ISTEXT({anchor}),ISNUMBER(SEARCH("MARGINAL",{anchor})))'],
            fill=AMBER_FILL, font=AMBER_FONT))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND(ISTEXT({anchor}),LEFT({anchor},3)="OK ")'], font=GREEN_FONT))


def main():
    wb = Workbook()
    wb.remove(wb.active)
    for mod_name in SHEET_MODULES:
        mod = importlib.import_module(f"suite.sheets.{mod_name}")
        mod.build(wb)
        print(f"  built {mod_name}")
    from suite.sheets import cover
    cover.build(wb)
    for ws in wb.worksheets:
        add_status_formatting(ws)
        # protect formulas; input cells are created unlocked (no password -
        # Review > Unprotect Sheet to edit structure)
        ws.protection.sheet = True
        ws.protection.formatColumns = False
        ws.protection.formatRows = False
    wb.save(OUT)
    print(f"Saved {OUT} ({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
