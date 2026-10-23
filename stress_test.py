"""Ten-well stress test: patch the built workbook into 10 different wells,
recalculate everything, and assert hand-independent physics.

Usage:  python stress_test.py         (expects petroleum_engineering_suite.xlsx)

Per scenario:
  - writes Well Data, the Directional Survey table, Kill Sheet kick data and
    string/annulus lengths, plus scenario-specific sheet inputs
  - recalculates the whole workbook with the `formulas` package
  - asserts: zero unexpected Excel errors; kill-sheet numbers (KMW/ICP/FCP/
    MAASP/volumes/schedule endpoints & midpoint) against an independent
    python model; survey TVD against python minimum curvature; torque & drag
    hookloads against an independent Johancsik march; scenario-specific
    checks (ESP, rod pump, chemistry, subsea, tubing stress, log, completion
    design); and expected status texts.
"""
import math
import os
import sys
import tempfile

import formulas
import openpyxl

BASE = "petroleum_engineering_suite.xlsx"
SCRATCH = tempfile.gettempdir()

ERRS = ("#DIV/0!", "#VALUE!", "#NAME?", "#REF!", "#NUM!", "#N/A", "#NULL!")

# geometry constants shared with the workbook defaults
CAP_DP = 4.276 ** 2 / 1029.4
CAP_HW = 3.0 ** 2 / 1029.4
CAP_DC = 2.813 ** 2 / 1029.4
SPR = 0.102
HOLE, CSG_ID = 8.5, 8.681
DPWT, BHAWT, BHALEN = 20.9, 92.0, 300
MUC, MUO, WOB, BITTQ, TJOD, BLOCK = 0.25, 0.30, 15000, 4000, 6.625, 85.0
TBG_OD, TBG_ID = 3.5, 2.992


# ---------------------------------------------------------------- helpers
def vertical(md, step=1000):
    pts = [(d, 0, 0) for d in range(0, md, step)]
    pts.append((md, 0, 0))
    return pts


def min_curv(stations):
    """[(md, tvd)] by minimum curvature."""
    out = [(stations[0][0], 0.0)]
    tvd = 0.0
    for i in range(1, len(stations)):
        md1, i1, a1 = stations[i - 1]
        md2, i2, a2 = stations[i]
        dmd = md2 - md1
        r1, r2 = math.radians(i1), math.radians(i2)
        cosdl = (math.cos(r1) * math.cos(r2)
                 + math.sin(r1) * math.sin(r2) * math.cos(math.radians(a2 - a1)))
        dl = math.acos(max(-1.0, min(1.0, cosdl)))
        rf = 1.0 if dl == 0 else 2 / dl * math.tan(dl / 2)
        tvd += dmd / 2 * (math.cos(r1) + math.cos(r2)) * rf
        out.append((md2, tvd))
    return out


def interp(prof, md):
    for i in range(1, len(prof)):
        if prof[i][0] >= md:
            m0, t0 = prof[i - 1]
            m1, t1 = prof[i]
            return t0 + (md - m0) * (t1 - t0) / (m1 - m0)
    return prof[-1][1]


def td_march(stations, mw, shoe, case):
    bf = (65.5 - mw) / 65.5
    mdbit = stations[-1][0]
    T = 0.0 if case == "pu" else -WOB
    M = BITTQ
    for i in range(len(stations) - 2, -1, -1):
        md1, i1, a1 = stations[i]
        md2, i2, a2 = stations[i + 1]
        dmd = md2 - md1
        th = math.radians((i1 + i2) / 2)
        w = bf * (BHAWT if md1 + dmd / 2 > mdbit - BHALEN else DPWT)
        mu = MUC if md1 + dmd / 2 <= shoe else MUO
        N = math.hypot(T * math.radians(a2 - a1) * math.sin(th),
                       T * math.radians(i2 - i1) + w * dmd * math.sin(th))
        if case == "pu":
            T += w * dmd * math.cos(th) + mu * N
        elif case == "so":
            T += w * dmd * math.cos(th) - mu * N
        else:
            M += mu * N * TJOD / 24
            T += w * dmd * math.cos(th)
    return T, M


def ceil10(x):
    # mirrors ROUNDUP(ROUND(x*10,6),0)/10
    return math.ceil(round(x * 10, 6)) / 10


def rceil(x):
    # mirrors ROUNDUP(ROUND(x,6),0)
    return math.ceil(round(x, 6))


class Scenario:
    def __init__(self, name, survey, omw, shoe_md, lot, sidpp, sicp, scr,
                 gain=10, extra_patches=(), extra_checks=(), extra_texts=(),
                 resp=None):
        self.name = name
        self.survey = survey
        self.omw, self.shoe_md, self.lot = omw, shoe_md, lot
        self.sidpp, self.sicp, self.scr, self.gain = sidpp, sicp, scr, gain
        self.extra_patches = list(extra_patches)
        self.extra_checks = list(extra_checks)
        self.extra_texts = list(extra_texts)
        self.md = survey[-1][0]
        self.prof = min_curv(survey)
        self.tvd = interp(self.prof, self.md)
        self.shoe_tvd = interp(self.prof, shoe_md)
        self.resp = resp if resp is not None else sidpp + 0.052 * omw * self.tvd

    # ------------------------------------------------------------ expected
    def kill_expected(self):
        s = self
        vstr = (CAP_DP * (s.md - 800) + CAP_HW * 500 + CAP_DC * 300)
        stb = rceil(vstr / SPR)
        kmwx = s.omw + s.sidpp / (0.052 * s.tvd)
        kmw = ceil10(kmwx)
        icp = s.scr + s.sidpp
        fcp = s.scr * kmw / s.omw
        checks = [
            ("Kill Sheet", "Drill string volume", vstr, 1e-4),
            ("Kill Sheet", "Strokes surface to bit", stb, 1e-9),
            ("Kill Sheet", "Kill mud weight (exact)", kmwx, 1e-6),
            ("Kill Sheet", "Kill mud weight (rounded up)", kmw, 1e-9),
            ("Kill Sheet", "Initial circulating pressure (ICP)", icp, 1e-9),
            ("Kill Sheet", "Final circulating pressure (FCP)", fcp, 1e-6),
            ("Kill Sheet", "MAASP (initial, with current mud)",
             rceil((s.lot - s.omw) * 0.052 * s.shoe_tvd), 1e-9),
            ("Kill Sheet", "Formation pressure",
             s.sidpp + 0.052 * s.omw * s.tvd, 1e-6),
            ("Kill Sheet", "Pressure drop per 100 strokes (vertical well)",
             math.floor(round((icp - fcp) / stb * 100, 6)), 1e-9),
        ]
        # schedule midpoint (50%) and endpoint (100%)
        sched = []
        for s_f in (0.5, 1.0):
            vt = s_f * vstr
            # locate the front in the tapered string (DP then HWDP then DC)
            if vt <= CAP_DP * (s.md - 800):
                fmd = vt / CAP_DP
            elif vt <= CAP_DP * (s.md - 800) + CAP_HW * 500:
                fmd = (s.md - 800) + (vt - CAP_DP * (s.md - 800)) / CAP_HW
            else:
                fmd = (s.md - 300) + (vt - CAP_DP * (s.md - 800)
                                      - CAP_HW * 500) / CAP_DC
            ftvd = interp(self.prof, fmd)
            p = round(s.sidpp - 0.052 * (kmw - s.omw) * ftvd + s.scr
                      + (fcp - s.scr) * s_f)
            sched.append((s_f, fmd, ftvd, p))
        return checks, sched

    def td_expected(self):
        s = self
        pu, _ = td_march(s.survey, s.omw, s.shoe_md, "pu")
        so, _ = td_march(s.survey, s.omw, s.shoe_md, "so")
        _, tq = td_march(s.survey, s.omw, s.shoe_md, "rot")
        return [
            ("Torque & Drag", "Pick-up hookload", pu / 1000 + BLOCK, 2e-3),
            ("Torque & Drag", "Slack-off hookload", so / 1000 + BLOCK, 2e-3),
            ("Torque & Drag", "Surface torque (rotating)", tq, 2e-3),
        ]


# ---------------------------------------------------------------- patching
def find_row(ws, label):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == label:
            return row[0].row
    raise KeyError(f"{ws.title}!{label}")


def apply(wb, scn):
    wd = wb["Well Data"]
    for label, val in {
        "Current mud / fluid weight": scn.omw,
        "Bit / TD measured depth": scn.md,
        "Bit / TD true vertical depth": scn.tvd,
        "Casing shoe MD": scn.shoe_md,
        "Casing shoe TVD": scn.shoe_tvd,
        "LOT / FIT equivalent mud weight": scn.lot,
        "Slow circulating rate pressure (SCR)": scn.scr,
        "Reservoir / formation pressure": scn.resp,
    }.items():
        wd.cell(row=find_row(wd, label), column=2, value=val)
    sv = wb["Directional Survey"]
    for i in range(100):
        r = 7 + i
        md, inc, azi = scn.survey[i] if i < len(scn.survey) else (None, None, None)
        # NOTE: ws.cell(..., value=None) does NOT clear a cell - assign .value
        sv.cell(row=r, column=1).value = md
        sv.cell(row=r, column=2).value = inc
        sv.cell(row=r, column=3).value = azi
    ks = wb["Kill Sheet"]
    for label, val in {
        "Shut-in drill pipe pressure (SIDPP)": scn.sidpp,
        "Shut-in casing pressure (SICP)": scn.sicp,
        "Pit gain": scn.gain,
    }.items():
        ks.cell(row=find_row(ks, label), column=2, value=val)
    for a_val, val in {
        "DP 5 19.50#": scn.md - 800,
        "DP x open hole": scn.md - scn.shoe_md - 300,
        "DP x casing": scn.shoe_md,
    }.items():
        ks.cell(row=find_row(ks, a_val), column=4, value=val)
    for sheet, label, val in scn.extra_patches:
        ws = wb[sheet]
        ws.cell(row=find_row(ws, label), column=2, value=val)


# ---------------------------------------------------------------- scenarios
def build_scenarios():
    scns = []

    scns.append(Scenario("1 vertical shallow gas", vertical(6000),
                         omw=9.2, shoe_md=2500, lot=12.5,
                         sidpp=300, sicp=450, scr=600))

    hz = ([(0, 0, 0)] + [(d, 0, 0) for d in range(1000, 6001, 1000)]
          + [(6000 + 100 * i, 8 * i, 30) for i in range(1, 12)]
          + [(7125, 90, 30)]
          + [(d, 90, 30) for d in range(7500, 12001, 500)])
    scns.append(Scenario("2 horizontal oil deep", hz,
                         omw=11.0, shoe_md=7125, lot=14.0,
                         sidpp=400, sicp=550, scr=800, gain=20))

    sw = ([(d, 0, 0) for d in range(0, 2001, 500)]
          + [(2000 + 100 * i, 4 * i, 120) for i in range(1, 11)]
          + [(d, 40, 120) for d in range(3500, 8001, 500)]
          + [(8000 + 100 * i, 40 - 4 * i, 120) for i in range(1, 11)]
          + [(9500, 0, 120), (10000, 0, 120)])
    scns.append(Scenario("3 S-well injector", sw,
                         omw=9.8, shoe_md=8000, lot=13.1,
                         sidpp=150, sicp=200, scr=700, gain=5))

    scns.append(Scenario(
        "4 HPHT vertical", vertical(16000, 2000),
        omw=17.5, shoe_md=14000, lot=19.2, sidpp=800, sicp=950, scr=1200,
        extra_patches=[("Well Data", "Perforation / reservoir TVD", 15500),
                       ("Tubing Stress Analysis", "Tubing surface pressure (B)", 12000)],
        extra_texts=[("Tubing Stress Analysis", "Triaxial status (B)",
                      "CHECK: EXCEEDS yield - redesign")]))

    scns.append(Scenario(
        "5 subsea deepwater", vertical(15000),
        omw=12.0, shoe_md=11000, lot=14.5, sidpp=350, sicp=500, scr=900,
        extra_patches=[("Kill Sheet - Subsea", "Water depth", 5000),
                       ("Kill Sheet - Subsea", "Choke line length", 5080),
                       ("Kill Sheet - Subsea",
                        "Choke line friction (CLF) at kill rate", 400),
                       ("Kill Sheet - Subsea",
                        "Shut-in drill pipe pressure (SIDPP)", 350),
                       ("Kill Sheet - Subsea",
                        "Shut-in casing pressure (SICP)", 500)],
        extra_checks=[("Kill Sheet - Subsea",
                       "Dynamic MAASP (circulating up choke line)", 1030.0, 1e-9)]))

    scns.append(Scenario(
        "6 sour gas", vertical(11000),
        omw=13.0, shoe_md=9000, lot=16.0, sidpp=600, sicp=750, scr=850,
        resp=6000,
        extra_patches=[("Production Chemistry", "CO2 content", 3.5),
                       ("Production Chemistry", "H2S content", 0.01)],
        extra_checks=[("Production Chemistry", "CO2 partial pressure (pCO2)",
                       210.0, 1e-6),
                      ("Production Chemistry", "H2S partial pressure (pH2S)",
                       0.6, 1e-6)]))

    rp = Scenario("7 heavy oil rod pump", vertical(4000),
                  omw=8.9, shoe_md=3500, lot=11.8,
                  sidpp=50, sicp=80, scr=400, gain=3,
                  extra_patches=[("Rod Pump Design", "Plunger diameter (D)", 1.75),
                                 ("Rod Pump Design", "Stroke length (S)", 100),
                                 ("Rod Pump Design", "Pumping speed (N)", 6),
                                 ("Rod Pump Design", "Pump depth (L)", 3800),
                                 ("Rod Pump Design", "Fluid specific gravity", 0.95),
                                 ("Rod Pump Design", "Taper 1 length", 1500),
                                 ("Rod Pump Design", "Taper 2 length", 2300),
                                 ("Rod Pump Design",
                                  "Net lift (H, fluid level depth)", 3000)])
    w1 = 2.904 * 0.875 ** 2 * 1500
    w2 = 2.904 * 0.75 ** 2 * 2300
    wr = w1 + w2
    wrf = wr * (1 - 0.128 * 0.95)
    fo = 0.340 * 0.95 * 1.75 ** 2 * 3000
    alpha = 100 * 36 / 70500
    rp.extra_checks += [
        ("Rod Pump Design", "Pump displacement (100% eff)",
         0.1166 * 100 * 6 * 1.75 ** 2, 1e-6),
        ("Rod Pump Design", "Rod weight in air (Wr)", wr, 1e-6),
        ("Rod Pump Design", "Peak polished rod load (PPRL)",
         fo + wrf * (1 + alpha), 1e-6),
        ("Rod Pump Design", "Minimum polished rod load (MPRL)",
         wrf * (1 - alpha), 1e-6),
    ]
    scns.append(rp)

    scns.append(Scenario(
        "8 gas lift well", vertical(9000),
        omw=9.0, shoe_md=8500, lot=13.0, sidpp=200, sicp=300, scr=650,
        extra_patches=[("Gas Lift Design",
                        "Kick-off / injection surface pressure (Pko)", 1200),
                       ("Gas Lift Design", "Wellhead flowing pressure (Pwh)", 120),
                       ("Gas Lift Design", "Load (kill) fluid gradient (Gs)", 0.45)]))

    esp = Scenario("9 ESP redesign", vertical(7000),
                   omw=9.5, shoe_md=6200, lot=13.5,
                   sidpp=250, sicp=350, scr=700,
                   extra_patches=[("ESP Design", "PIP for design (override)", 1100),
                                  ("ESP Design", "Composite SG for design", 1.05),
                                  ("ESP Design", "Pump setting depth (TVD)", 6500),
                                  ("ESP Design", "Head per stage (catalog, at BEP)", 22)])
    fop = 1100 * 2.31 / 1.05
    tdh = (6500 - fop) + 150 * 2.31 / 1.05 + 20 * 6.5
    esp.extra_checks += [
        ("ESP Design", "Total dynamic head (TDH)", tdh, 1e-4),
        ("ESP Design", "Stages required", rceil(tdh / 22), 1e-9),
    ]
    scns.append(esp)

    cd = Scenario("10 frac pack well", vertical(8200),
                  omw=10.2, shoe_md=7800, lot=14.2,
                  sidpp=180, sicp=260, scr=600,
                  extra_patches=[("Completion Design", "d50 grain size (median)", 90),
                                 ("Completion Design", "C factor", 125),
                                 ("Completion Design", "Mixture density", 45),
                                 ("Log Interpretation",
                                  "True formation resistivity (Rt)", 5),
                                 ("Log Interpretation", "Porosity for Sw (PHI)", 0.18)])
    cd.extra_checks += [
        ("Completion Design", "Gravel median size (Saucier)", 540.0, 1e-6),
        ("Completion Design", "Erosional velocity", 125 / math.sqrt(45), 1e-4),
        ("Log Interpretation", "Water saturation (Archie)",
         math.sqrt(0.05 / (0.18 ** 2 * 5)), 1e-4),
    ]
    cd.extra_texts += [("Completion Design", "Commercial gravel selection", "30/50")]
    scns.append(cd)
    return scns


# ---------------------------------------------------------------- runner
def na_whitelist(wb):
    cells = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "NA(" in cell.value:
                    cells.add(f"{ws.title.upper()}'!{cell.coordinate}")
    return cells


def run_scenario(scn, idx):
    wb = openpyxl.load_workbook(BASE)
    apply(wb, scn)
    # locate labels/schedule rows before evaluation
    addr = {}
    checks, sched = scn.kill_expected()
    all_checks = checks + scn.td_expected() + scn.extra_checks
    for sheet, label, _, _ in all_checks:
        addr[(sheet, label)] = f"B{find_row(wb[sheet], label)}"
    texts = [("Kill Sheet", "Length check vs bit MD", "OK - matches bit MD")]
    texts += scn.extra_texts
    for sheet, label, _ in texts:
        addr[(sheet, label)] = f"B{find_row(wb[sheet], label)}"
    ks_hdr = find_row(wb["Kill Sheet"], "% strokes to bit")
    ok_na = na_whitelist(wb)

    path = os.path.join(SCRATCH, f"well_{idx}.xlsx")
    wb.save(path)
    xl = formulas.ExcelModel().loads(path).finish()
    sol = xl.calculate()
    book = path.replace("\\", "/").split("/")[-1]

    def get(sheet, a):
        return sol[f"'[{book}]{sheet.upper()}'!{a}"].value[0, 0]

    fails = []
    n_err = 0
    for key, val in sol.items():
        if ":" in key.split("!")[-1]:
            continue  # range aggregations; member cells are scanned individually
        try:
            v = val.value[0, 0]
        except Exception:
            continue
        if isinstance(v, str) and v.strip().upper() in ERRS:
            if v.strip().upper() == "#N/A" and any(key.endswith(k) for k in ok_na):
                continue
            n_err += 1
            if n_err <= 5:
                fails.append(f"error cell {key} = {v}")
    if n_err > 5:
        fails.append(f"... {n_err} error cells total")

    for sheet, label, exp, tol in all_checks:
        got = get(sheet, addr[(sheet, label)])
        if not isinstance(got, (int, float)) or abs(got - exp) > tol * max(1, abs(exp)):
            fails.append(f"{sheet}!{label}: got {got!r}, expected {exp:.5g}")
    for sheet, label, exp in texts:
        got = get(sheet, addr[(sheet, label)])
        if not (isinstance(got, str) and exp in got):
            fails.append(f"{sheet}!{label}: got {got!r}, expected text ~ {exp!r}")
    # schedule rows: 0%..100% in 5% steps -> 50% = row hdr+11, 100% = hdr+21
    for s_f, fmd, ftvd, p in sched:
        r = ks_hdr + 1 + int(s_f * 20)
        for col, exp, tol in (("C", fmd, 1.0), ("D", ftvd, 1.5), ("E", p, 1.1)):
            got = get("Kill Sheet", f"{col}{r}")
            if not isinstance(got, (int, float)) or abs(got - exp) > tol:
                fails.append(f"schedule {int(s_f*100)}% col {col}: got {got!r}, "
                             f"expected {exp:.1f}")
    return fails


def main():
    total_fail = 0
    for idx, scn in enumerate(build_scenarios(), start=1):
        fails = run_scenario(scn, idx)
        status = "PASS" if not fails else "FAIL"
        print(f"[{status}] {scn.name}  (MD {scn.md}, TVD {scn.tvd:.0f})")
        for f in fails:
            print("        ", f)
        total_fail += len(fails)
    print(f"\n{'ALL 10 WELLS PASS' if not total_fail else str(total_fail) + ' FAILURES'}")
    sys.exit(1 if total_fail else 0)


if __name__ == "__main__":
    main()
